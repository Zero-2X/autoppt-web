"""Paid add-ons: speaker notes and anticipated defense questions.

Planning produces the content (``speaker_notes`` per slide and a top-level
``anticipated_questions`` list in ``deck-plan.json``) when the order includes
the add-on; the Worker checks it is there.  Packaging turns it into
``output/speaker-notes.md`` / ``output/anticipated-questions.md`` and writes the
notes into the PPTX's speaker-notes pages, so the user never pays for an add-on
that was not delivered.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

MIN_QUESTIONS = 8


def addon_errors(plan: dict[str, Any], addons: list[str]) -> list[str]:
    errors = []
    if "speaker_notes" in addons:
        missing = [s["slide_id"] for s in plan["slides"] if not str(s.get("speaker_notes", "")).strip()]
        if missing:
            errors.append(f"the order includes speaker notes; write speaker_notes for {', '.join(missing)}")
    if "anticipated_questions" in addons:
        questions = [q for q in plan.get("anticipated_questions", []) if str(q.get("question", "")).strip()]
        if len(questions) < MIN_QUESTIONS:
            errors.append(f"the order includes anticipated questions; write at least {MIN_QUESTIONS} in anticipated_questions "
                          "(question, answer_outline, evidence_refs)")
    return errors


def write_addon_files(job_dir: Path, plan: dict[str, Any], addons: list[str]) -> list[Path]:
    out = Path(job_dir) / "output"
    out.mkdir(parents=True, exist_ok=True)
    files = []
    if "speaker_notes" in addons:
        lines = [f"# {plan.get('deck', {}).get('title', '')} · 讲稿", ""]
        for slide in plan["slides"]:
            lines += [f"## {slide['slide_id']} · {slide.get('headline', '')}", "", str(slide.get("speaker_notes", "")).strip(), ""]
        path = out / "speaker-notes.md"
        path.write_text("\n".join(lines), encoding="utf-8")
        files.append(path)
    if "anticipated_questions" in addons:
        lines = [f"# {plan.get('deck', {}).get('title', '')} · 预测问题", ""]
        for index, item in enumerate(plan.get("anticipated_questions", []), 1):
            lines += [f"## {index}. {item.get('question', '')}", "", f"回答要点：{item.get('answer_outline', '')}"]
            if item.get("evidence_refs"):
                lines.append(f"依据：{', '.join(item['evidence_refs'])}")
            lines.append("")
        path = out / "anticipated-questions.md"
        path.write_text("\n".join(lines), encoding="utf-8")
        files.append(path)
    return files


def add_notes_to_deck(pptx: Path, plan: dict[str, Any]) -> int:
    """Write each slide's speaker notes into the deck's notes pages; returns slides annotated."""
    from pptx import Presentation

    presentation = Presentation(str(pptx))
    notes = {slide["slide_id"]: str(slide.get("speaker_notes", "")).strip() for slide in plan["slides"]}
    written = 0
    for index, slide in enumerate(presentation.slides, 1):
        text = notes.get(f"S{index:02d}", "")
        if text:
            slide.notes_slide.notes_text_frame.text = text
            written += 1
    presentation.save(str(pptx))
    return written
