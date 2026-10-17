"""PPTX: per-slide text, notes, tables, chart data and pictures (python-pptx)."""

from __future__ import annotations

from pathlib import Path
from typing import Any, Iterator

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

from autoppt_service.parsers.archive import inspect_ooxml
from autoppt_service.parsers.base import (
    IntakeLimits,
    ParsedSource,
    SourceFile,
    SourceRejected,
    image_asset,
    make_item,
    normalize_text,
    table_item,
    table_record,
    text_items,
)


def _walk(shapes) -> Iterator[Any]:
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _walk(shape.shapes)
        else:
            yield shape


def _chart_rows(chart) -> list[list[Any]]:
    plot = chart.plots[0] if len(chart.plots) else None
    if plot is None:
        return []
    try:  # XY and bubble charts have no category axis
        categories = [str(value) for value in plot.categories]
        series = list(plot.series)
    except Exception:
        return []
    rows: list[list[Any]] = [["category", *[s.name for s in series]]]
    for index, category in enumerate(categories):
        rows.append([category, *[s.values[index] if index < len(s.values) else None for s in series]])
    return rows


def parse_pptx(source: SourceFile, work_dir: Path, limits: IntakeLimits) -> ParsedSource:
    inspect_ooxml(source.path, limits)
    try:
        presentation = Presentation(str(source.path))
    except Exception as exc:
        raise SourceRejected("pptx_unreadable", f"无法读取 PPTX：{exc}") from exc
    slides = list(presentation.slides)
    if len(slides) > limits.max_pptx_slides:
        raise SourceRejected("pptx_too_many_slides", f"PPTX 页数 {len(slides)} 超过上限 {limits.max_pptx_slides}")

    result = ParsedSource(parser="pptx")
    media_dir = work_dir / "media" / source.file_id
    image_count = 0
    for number, slide in enumerate(slides, 1):
        title = ""
        if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            title = normalize_text(slide.shapes.title.text_frame.text)
        texts: list[str] = []
        for shape in _walk(slide.shapes):
            location = f"slide {number}"
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    line = normalize_text("".join(run.text for run in paragraph.runs))
                    if line and line != title:
                        texts.append(line)
            if getattr(shape, "has_table", False) and shape.has_table:
                rows = [[normalize_text(cell.text) for cell in row.cells] for row in shape.table.rows]
                record = table_record(source, f"s{number:03d}-t{len(result.tables) + 1:02d}", rows, location=location)
                result.tables.append(record)
                result.items.append(table_item(source, record, page=number))
            if getattr(shape, "has_chart", False) and shape.has_chart:
                rows = _chart_rows(shape.chart)
                if rows:
                    record = table_record(
                        source,
                        f"s{number:03d}-chart{len(result.tables) + 1:02d}",
                        rows,
                        location=f"{location} chart",
                        extra={"from_chart": str(shape.chart.chart_type)},
                    )
                    result.tables.append(record)
                    result.items.append(table_item(source, record, page=number))
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and image_count < limits.max_embedded_images:
                try:
                    image = shape.image
                except Exception:  # linked (not embedded) pictures have no blob
                    result.warnings.append(f"第 {number} 页有无法读取的图片（可能是外部链接）")
                    continue
                media_dir.mkdir(parents=True, exist_ok=True)
                suffix = f"s{number:03d}-img{image_count + 1:02d}"
                target = media_dir / f"{suffix}.{image.ext}"
                target.write_bytes(image.blob)
                width, height = image.size
                image_count += 1
                result.assets.append(
                    image_asset(
                        source, suffix, target, origin=location, width=width, height=height, image_format=image.ext.upper()
                    )
                )
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        body = "\n".join(([title] if title else []) + texts)
        if body:
            result.items.extend(
                text_items(
                    source,
                    f"s{number:03d}",
                    body,
                    limits,
                    title=title or f"{source.name} 第 {number} 页",
                    location=f"slide {number}",
                    page=number,
                    chapter=title or None,
                )
            )
        if notes:
            result.items.append(
                make_item(
                    source,
                    f"s{number:03d}-notes",
                    item_type="speaker_notes",
                    text=notes[: limits.max_item_chars],
                    title=f"{source.name} 第 {number} 页备注",
                    location=f"slide {number} notes",
                    page=number,
                )
            )
        if not body and source.role == "vectorize_source":
            result.warnings.append(f"第 {number} 页没有可抽取文字（可能是整页图片）")

    result.stats = {
        "slides": len(slides),
        "slide_width_emu": int(presentation.slide_width or Emu(0)),
        "slide_height_emu": int(presentation.slide_height or Emu(0)),
        "tables": len(result.tables),
        "images": image_count,
    }
    return result
