"""XLSX / CSV: chart data tables with a numeric-column hint."""

from __future__ import annotations

import csv
import io
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

from autoppt_service.parsers.archive import inspect_ooxml
from autoppt_service.parsers.base import (
    IntakeLimits,
    ParsedSource,
    SourceFile,
    SourceRejected,
    decode_text,
    table_item,
    table_record,
)


def _trim(rows: Iterable[Iterable[Any]], limits: IntakeLimits) -> tuple[list[list[Any]], bool]:
    kept: list[list[Any]] = []
    truncated = False
    for row in rows:
        values = list(row)
        if len(values) > limits.max_sheet_cols:
            values = values[: limits.max_sheet_cols]
            truncated = True
        while values and (values[-1] is None or values[-1] == ""):
            values.pop()
        if not values and not kept:
            continue  # leading blank rows
        if len(kept) > limits.max_sheet_rows:  # header + max rows
            truncated = True
            break
        kept.append(values)
    while kept and not kept[-1]:
        kept.pop()
    return kept, truncated


def _add(
    result: ParsedSource,
    source: SourceFile,
    limits: IntakeLimits,
    name: str,
    rows: list[list[Any]],
    truncated: bool,
    location: str,
    extra: dict[str, Any] | None = None,
) -> None:
    if not rows:
        result.warnings.append(f"{location} 为空")
        return
    record = table_record(source, name, rows, location=location, truncated=truncated, extra=extra)
    result.tables.append(record)
    result.items.append(table_item(source, record, title=f"{source.name} {location}"))
    if truncated:
        result.warnings.append(
            f"{location} 超过读取上限（{limits.max_sheet_rows} 行 / {limits.max_sheet_cols} 列），已截断"
        )


def parse_xlsx(source: SourceFile, work_dir: Path, limits: IntakeLimits) -> ParsedSource:
    inspect_ooxml(source.path, limits)
    try:
        workbook = load_workbook(str(source.path), read_only=True, data_only=True)
    except Exception as exc:
        raise SourceRejected("xlsx_unreadable", f"无法读取 XLSX：{exc}") from exc
    result = ParsedSource(parser="xlsx")
    try:
        for index, sheet in enumerate(workbook.worksheets, 1):
            rows, truncated = _trim(sheet.iter_rows(values_only=True), limits)
            location = f"sheet {sheet.title}"
            _add(result, source, limits, f"sheet{index:02d}", rows, truncated, location, {"sheet": sheet.title})
    finally:
        workbook.close()
    result.stats = {"sheets": len(workbook.sheetnames), "tables": len(result.tables)}
    return result


def parse_csv(source: SourceFile, work_dir: Path, limits: IntakeLimits) -> ParsedSource:
    text, encoding = decode_text(source.path.read_bytes())
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows, truncated = _trim(csv.reader(io.StringIO(text), dialect), limits)
    result = ParsedSource(parser="csv")
    _add(result, source, limits, "table01", rows, truncated, "csv")
    result.stats = {"encoding": encoding, "tables": len(result.tables)}
    if encoding == "latin-1":
        result.warnings.append("无法识别文本编码，按 latin-1 读取，可能有乱码")
    return result
