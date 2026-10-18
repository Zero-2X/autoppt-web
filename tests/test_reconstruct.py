"""Real reconstruction phase with a fake Codex agent that writes PPTX files, and the Worker's audit."""

from __future__ import annotations

import json
import re
import sys
from pathlib import Path

import pytest
from PIL import Image

from autoppt_service.checkpoint import pending_checkpoint, submit_response
from autoppt_service.jobfiles import read_json, service_path, write_json_atomic
from autoppt_service.reconstruct import audit_pptx
from autoppt_service.run_job import run_phase
from autoppt_service.status import load_job_state
from autoppt_service.worker_template import build_template

FAKE_BUILDER = r'''
import json, os, sys, uuid
from pathlib import Path
from pptx import Presentation
from pptx.util import Emu, Pt, Inches

prompt = sys.stdin.buffer.read().decode("utf-8")
home = Path(os.environ["CODEX_HOME"])
counter = home / "calls.txt"
calls = int(counter.read_text()) + 1 if counter.exists() else 1
counter.write_text(str(calls))
(home / f"prompt-{calls}.txt").write_text(prompt, encoding="utf-8")
mode = os.environ.get("FAKE_BUILD_MODE", "ok")
print(json.dumps({"type": "thread.started", "thread_id": str(uuid.uuid4())}))
if mode == "missing_then_limit" and calls >= 2:
    print(json.dumps({"type": "error", "message": "You've hit your usage limit. Try again at 3:13 PM."}))
    print(json.dumps({"type": "turn.failed", "error": {"message": "usage limit"}}))
    sys.exit(1)
job = Path.cwd()
plan = json.loads((job / "workspace/presentation_handoff/deck-plan.json").read_text(encoding="utf-8"))
deck = Presentation()
deck.slide_width, deck.slide_height = Emu(12192000), Emu(6858000)
for slide in plan["slides"]:
    page = deck.slides.add_slide(deck.slide_layouts[6])
    if mode == "picture":
        page.shapes.add_picture(str(job / "assets/slides" / f"{slide['slide_id']}.png"), 0, 0, deck.slide_width, deck.slide_height)
        continue
    lines = list(slide["exact_text"])
    if mode in ("missing_then_ok", "missing_then_limit") and "is not native text" not in prompt and slide["slide_id"] == "S02":
        lines = lines[:-1]
    for index, line in enumerate(lines):
        box = page.shapes.add_textbox(Inches(0.6), Inches(0.5 + 0.8 * index), Inches(12), Inches(0.7))
        box.text_frame.text = line
        box.text_frame.paragraphs[0].runs[0].font.name = "Microsoft YaHei"
        box.text_frame.paragraphs[0].runs[0].font.size = Pt(24)
    if "<<<USER_REPLY" in prompt and slide["slide_id"] == "S01":
        note = page.shapes.add_textbox(Inches(0.6), Inches(6.5), Inches(12), Inches(0.5))
        note.text_frame.text = "已按意见调整"
(job / "output").mkdir(exist_ok=True)
deck.save(job / "output/native-deck.pptx")
report = {"route": "measured", "slides": [], "disclosures": ["test"], "learning_context": {"files_read": ["x"]}}
(job / "workspace/reconstruct").mkdir(parents=True, exist_ok=True)
(job / "workspace/reconstruct/reconstruction-report.json").write_text(json.dumps(report), encoding="utf-8")
(job / "service/status.json").write_text(json.dumps({"schema_version": "autoppt-agent-status-v1", "phase": "reconstruct",
    "state": "done", "updated_at": "2027-01-01T00:00:00Z"}), encoding="utf-8")
print(json.dumps({"type": "turn.completed", "usage": {"input_tokens": 5000, "output_tokens": 900}}))
'''


def fake_renderer(pptx: Path, out_dir: Path) -> list[Path]:
    from pptx import Presentation

    out_dir.mkdir(parents=True, exist_ok=True)
    paths = []
    for index in range(1, len(Presentation(str(pptx)).slides) + 1):
        target = out_dir / f"slide-{index}.png"
        Image.new("RGB", (1600, 900), (230, 235, 240)).save(target)
        paths.append(target)
    return paths


@pytest.fixture()
def job_at_reconstruct(tmp_path, monkeypatch):
    """A native job driven by the mock up to an approved CP2, with a reviewed deck plan."""
    from test_intake_and_run_job import make_request
    from test_worker_template import fake_repo
    from docx import Document

    uploads = tmp_path / "uploads"
    uploads.mkdir()
    document = Document()
    document.add_paragraph("城市内涝频发，现有预警只能给出区县级结果。")
    document.save(uploads / "paper.docx")
    request = make_request(files=[{"file_id": "f_01", "role": "primary_source", "name": "论文.docx"}])
    request["delivery"]["charts"] = "in_image"
    job = tmp_path / "job"
    run_phase(job, "intake", job_id="R1", request=request, uploads={"f_01": str(uploads / "paper.docx")})
    run_phase(job, "planning", job_id="R1", mock=True)
    cp = pending_checkpoint(job)
    submit_response(job, {"checkpoint_id": cp["checkpoint_id"], "decision": "approve", "responder": "user", "user_reply": "好"})
    run_phase(job, "masters", job_id="R1", mock=True)
    cp = pending_checkpoint(job)
    submit_response(job, {"checkpoint_id": cp["checkpoint_id"], "decision": "approve", "responder": "user", "user_reply": "好"})
    slides = read_json(job / "workspace/presentation_handoff/slide_brief.json")["slides"]
    write_json_atomic(job / "workspace/presentation_handoff/deck-plan.json", {"slides": [
        {"slide_id": s["slide_id"], "exact_text": [f"{s['slide_id']} 标题", f"{s['slide_id']} 第二行 47 分钟"]} for s in slides]})
    template = build_template(tmp_path / "templates", repo=fake_repo(tmp_path / "autoppt"), version="v1")
    script = tmp_path / "fake_builder.py"
    script.write_text(FAKE_BUILDER, encoding="utf-8")
    home = tmp_path / "home"
    home.mkdir()
    profiles = tmp_path / "profiles.json"
    profiles.write_text(json.dumps([{"name": "p", "codex_home": str(home)}]), encoding="utf-8")
    import autoppt_service.vectorize as vectorize

    monkeypatch.setattr(vectorize, "powerpoint_renderer", fake_renderer)
    plan_path = job / "workspace/presentation_handoff/deck-plan.json"

    def fake_ocr(image, out_json):  # renders show the planned text; masters show nothing OCR can read
        if "office-render" not in str(image):
            return {"lines": []}
        lines = [line for s in json.loads(plan_path.read_text(encoding="utf-8"))["slides"] for line in s["exact_text"]]
        return {"lines": [{"text": line} for line in lines]}

    monkeypatch.setattr(vectorize, "windows_ocr", fake_ocr)

    def reconstruct(mode: str = "ok"):
        monkeypatch.setenv("FAKE_BUILD_MODE", mode)
        return run_phase(job, "reconstruct", job_id="R1", template=template, profiles_path=profiles,
                         command=[sys.executable, str(script)])

    return {"job": job, "home": home, "reconstruct": reconstruct}


def test_reconstruct_audits_renders_and_opens_cp3(job_at_reconstruct):
    job = job_at_reconstruct["job"]
    result = job_at_reconstruct["reconstruct"]()
    assert result["job_state"] == "cp3_native_preview", result
    cp3 = pending_checkpoint(job)
    assert cp3["type"] == "native_preview"
    assert cp3["editability_scope"]["native"] == ["文字"] and cp3["editability_scope"]["counts"]["text_boxes"] > 0
    kinds = {p["kind"] for p in cp3["previews"]}
    assert {"side_by_side", "diff_heatmap", "slide"} <= kinds
    assert (job / "workspace/office-render/S01.png").exists()
    report = read_json(service_path(job, "reconstruct-report.json"))
    assert report["agent_report"]["route"] == "measured"
    agents = (job / "AGENTS.md").read_text(encoding="utf-8")
    assert "## 7. Phase: reconstruct" in agents and "{{" not in agents


def test_missing_native_text_is_sent_back_then_fixed(job_at_reconstruct):
    result = job_at_reconstruct["reconstruct"]("missing_then_ok")
    assert result["job_state"] == "cp3_native_preview"
    report = read_json(service_path(job_at_reconstruct["job"], "reconstruct-report.json"))
    assert len(report["attempts"]) == 2 and "S02: reviewed text is not native text" in report["attempts"][0]["errors"][0]
    assert "is not native text" in (job_at_reconstruct["home"] / "prompt-2.txt").read_text(encoding="utf-8")


def test_a_picture_of_the_master_is_never_accepted(job_at_reconstruct):
    result = job_at_reconstruct["reconstruct"]("picture")
    assert result["job_state"] == "blocked"
    assert "no native text" in load_job_state(job_at_reconstruct["job"])["blocked_reason"]


def test_cp3_revision_round_passes_user_words(job_at_reconstruct):
    job = job_at_reconstruct["job"]
    job_at_reconstruct["reconstruct"]()
    cp3 = pending_checkpoint(job)
    submit_response(job, {"checkpoint_id": cp3["checkpoint_id"], "decision": "revise", "responder": "user",
                          "user_reply": "第一页标题再大一点"})
    result = job_at_reconstruct["reconstruct"]()
    assert result["job_state"] == "cp3_native_preview" and pending_checkpoint(job)["round"] == 2
    prompts = sorted((job_at_reconstruct["home"]).glob("prompt-*.txt"))
    assert "第一页标题再大一点" in prompts[-1].read_text(encoding="utf-8")


def test_audit_rules(tmp_path):
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Emu, Inches

    image = tmp_path / "a.png"
    Image.new("RGB", (100, 60)).save(image)
    deck = Presentation()
    deck.slide_width, deck.slide_height = Emu(12192000), Emu(6858000)
    tiled = deck.slides.add_slide(deck.slide_layouts[6])
    box = tiled.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    box.text_frame.text = "标题"
    box.text_frame.paragraphs[0].runs[0].font.name = "Arial"
    for row in range(3):
        for col in range(3):
            tiled.shapes.add_picture(str(image), Emu(col * 4064000), Emu(row * 2286000), Emu(4064000), Emu(2286000))
    path = tmp_path / "deck.pptx"
    deck.save(path)
    errors, warnings, scope = audit_pptx(path, {"S01": ["标题"]}, slide_ids=["S01"], require_chart=True)
    assert errors[-1] == "the order has chart data but the deck has no native chart with an embedded workbook"
    assert "no native shapes" in errors[0]  # every slide is pictures plus text
    assert any("raster tiling" in w for w in warnings) and any("arial" in w for w in warnings)
    chart_slide = deck.slides.add_slide(deck.slide_layouts[6])
    data = CategoryChartData()
    data.categories = ["A", "B"]
    data.add_series("命中率", (0.61, 0.82))
    chart_slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(6), Inches(4), data)
    deck.save(path)
    errors, _, scope = audit_pptx(path, {"S01": ["标题"]}, slide_ids=["S01", "S02"], require_chart=True)
    assert errors == [] and scope["charts_with_workbook"] == 1 and "图表" in scope["native"]
    errors, _, _ = audit_pptx(path, {}, slide_ids=["S01"], require_chart=False)
    assert errors and "2 slides in the deck but 1 masters" in errors[0]


def test_stale_lettering_and_missing_shapes_fail_the_audit(tmp_path):
    from pptx import Presentation
    from pptx.util import Emu, Inches

    from autoppt_service.reconstruct import stale_lines

    assert stale_lines(["去掉管网约束后误报率上升到 0.26", "本文方法"], "误报率上升到 0.26") == ["去掉管网约束后误报率上升到 0.26"]
    assert stale_lines(["本文方法"], "完全不同的文字") == []
    background = tmp_path / "bg.png"
    Image.new("RGB", (160, 90), "white").save(background)
    deck = Presentation()
    deck.slide_width, deck.slide_height = Emu(12192000), Emu(6858000)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_picture(str(background), 0, 0, deck.slide_width, deck.slide_height)
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    box.text_frame.text = "对比与消融结果：本文方法精度更高"
    path = tmp_path / "deck.pptx"
    deck.save(path)

    def ocr(image, out_json):  # the background still carries the title lettering
        return {"lines": [{"text": "对 比 与 消 融 结 果 ： 本 文 方 法 精 度 更 高"}]}

    errors, _, _ = audit_pptx(path, {"S01": ["对比与消融结果：本文方法精度更高"]}, slide_ids=["S01"], require_chart=False,
                              ocr=ocr, work_dir=tmp_path / "w")
    assert any("still shows the old text" in e for e in errors)
    assert any("no native shapes" in e for e in errors)


def test_withdrawn_preview_sends_audit_errors_not_user_words(job_at_reconstruct):
    from autoppt_service.checkpoint import closed_checkpoints
    from autoppt_service.run_job import withdraw_native_preview

    job = job_at_reconstruct["job"]
    job_at_reconstruct["reconstruct"]()
    state = withdraw_native_preview(job, ["S01: the background picture still shows the old text"], reason="stale lettering")
    assert state["state"] == "reconstructing"
    checkpoint, response = closed_checkpoints(job)[-1]
    assert response["decision"] == "withdrawn" and response["responder"] == "system" and response["user_reply"] == ""
    result = job_at_reconstruct["reconstruct"]()
    assert result["job_state"] == "cp3_native_preview" and pending_checkpoint(job)["round"] == 2
    prompt = sorted(job_at_reconstruct["home"].glob("prompt-*.txt"))[-1].read_text(encoding="utf-8")
    assert "still shows the old text" in prompt and "<<<USER_REPLY" not in prompt


def test_render_visibility_finds_covered_and_lost_text(tmp_path):
    from autoppt_service.reconstruct import invisible_lines, render_visibility_errors

    assert invisible_lines(["去掉管网约束后误报率上升到 0.26"], "去掉管网约束后") == ["去掉管网约束后误报率上升到 0.26"]
    assert invisible_lines(["去掉管网约束后误报率上升到 0.26"], "去掉管网约束后误报率上升到0.26") == []
    renders = {"S07": tmp_path / "render.png"}
    masters = {"S07": tmp_path / "master.png"}
    for path in (*renders.values(), *masters.values()):
        Image.new("RGB", (10, 10)).save(path)

    def ocr(image, out_json):
        if Path(image).name == "master.png":
            return {"lines": [{"text": "对比与消融结果：本文方法精度更高"}, {"text": "去掉管网约束后误报率上升到 0.26"}]}
        return {"lines": [{"text": "方法 命中率 误报率"}, {"text": "去掉管网约束后"}]}

    errors = render_visibility_errors(renders, {"S07": ["去掉管网约束后误报率上升到 0.26"]}, ocr=ocr, masters=masters,
                                      work_dir=tmp_path / "w")
    assert "not visible in the PowerPoint render" in errors[0]
    assert "missing from the native render" in errors[1] and "对比与消融结果" in errors[1]


def test_usage_limit_defers_and_keeps_the_open_audit_issues(job_at_reconstruct):
    # Measured 2026-10-18: the last session hit the usage limit and the S01 finding was lost.
    job = job_at_reconstruct["job"]
    result = job_at_reconstruct["reconstruct"]("missing_then_limit")
    assert result.get("deferred") is True, result
    assert load_job_state(job)["state"] != "blocked"
    carried = read_json(service_path(job, "reconstruct-issues.json"))
    assert any("S02: reviewed text is not native text" in issue for issue in carried["issues"])
    assert read_json(service_path(job, "deferred.json"))["retry_after"]


def test_chart_axis_ids_powerpoint_cannot_read_are_renumbered(tmp_path):
    # Measured 2026-10-18: PowerPoint refused a deck whose chart used axis id 2226939960.
    import zipfile
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches
    from autoppt_service.reconstruct import normalize_chart_axis_ids

    deck = Presentation()
    data = CategoryChartData()
    data.categories = ["A", "B"]
    data.add_series("s", (1, 2))
    deck.slides.add_slide(deck.slide_layouts[6]).shapes.add_chart(
        XL_CHART_TYPE.BAR_CLUSTERED, 0, 0, Inches(4), Inches(3), data)
    path = tmp_path / "deck.pptx"
    deck.save(path)
    assert normalize_chart_axis_ids(path) == []  # python-pptx ids are already in range
    with zipfile.ZipFile(path) as package:
        items = [(i, package.read(i.filename)) for i in package.infolist()]
    with zipfile.ZipFile(path, "w") as package:
        for info, raw in items:
            if info.filename == "ppt/charts/chart1.xml":
                text = raw.decode("utf-8")
                for old in sorted(set(re.findall(r'axId val="(-?\d+)"', text))):
                    text = text.replace(f'val="{old}"', f'val="{2226939960 + len(old)}"')
                raw = text.encode("utf-8")
            package.writestr(info, raw)
    assert normalize_chart_axis_ids(path) == ["ppt/charts/chart1.xml"]
    text = zipfile.ZipFile(path).read("ppt/charts/chart1.xml").decode("utf-8")
    ids = [int(v) for v in re.findall(r'(?:axId|crossAx) val="(-?\d+)"', text)]
    assert ids and max(ids) <= 2**31 - 1
    axes = set(re.findall(r'<c:axId val="(\d+)"', text))
    assert set(re.findall(r'<c:crossAx val="(\d+)"', text)) <= axes  # cross references still match
    Presentation(str(path))  # still a valid package


def test_unopenable_slides_are_named(tmp_path):
    from pptx import Presentation
    from autoppt_service.reconstruct import locate_unopenable_slides

    deck = Presentation()
    for word in ("ok", "broken", "ok"):
        deck.slides.add_slide(deck.slide_layouts[6]).shapes.add_textbox(0, 0, 100, 100).text_frame.text = word
    deck.save(tmp_path / "deck.pptx")

    def renderer(pptx, out_dir):
        if "broken" in "".join(s.shapes[0].text_frame.text for s in Presentation(str(pptx)).slides):
            raise RuntimeError("PowerPoint could not open the file.")
        return []

    assert locate_unopenable_slides(tmp_path / "deck.pptx", renderer, ["S01", "S02", "S03"], tmp_path / "w") == ["S02"]


def test_resumed_job_rechecks_the_deck_before_starting_an_agent(job_at_reconstruct):
    job = job_at_reconstruct["job"]
    assert job_at_reconstruct["reconstruct"]("missing_then_limit").get("deferred") is True
    result = job_at_reconstruct["reconstruct"]("ok")  # the deck on disk still lacks a S02 line
    assert result["job_state"] == "cp3_native_preview", result
    attempts = read_json(service_path(job, "reconstruct-report.json"))["attempts"]
    assert attempts[0]["worker_recheck"] and "S02" in attempts[0]["errors"][0]
    assert attempts[1]["attempt"] == 1 and len(attempts) == 2
    assert not service_path(job, "reconstruct-issues.json").exists()


def test_small_visible_text_is_reread_on_an_upscaled_render(tmp_path):
    # Measured 2026-10-18 on S05/S08: small visible lines OCR skipped at 1x, and a legend fragment on the master.
    from autoppt_service.reconstruct import render_visibility_errors

    renders = {"S05": tmp_path / "render.png"}
    masters = {"S05": tmp_path / "master.png"}
    for path in (*renders.values(), *masters.values()):
        Image.new("RGB", (10, 10)).save(path)
    line = "输出街区积水风险等级（低、中、高）"

    def ocr(image, out_json):
        name = Path(image).name
        if name == "master.png":
            return {"lines": [{"text": line}, {"text": "0 准确率到 FI 值"}]}
        return {"lines": [{"text": "方法分三步"}] + ([{"text": line}] if name.endswith("-2x.png") else [])}

    assert render_visibility_errors(renders, {"S05": [line]}, ocr=ocr, masters=masters, work_dir=tmp_path / "w") == []
    assert (tmp_path / "w" / "S05-render-2x.png").exists()


def test_review_findings_are_never_settled_by_a_worker_recheck(job_at_reconstruct):
    # 2026-10-18: CP3 r3 was withdrawn for a white box behind the S01 title, which no audit can see.
    from autoppt_service.run_job import withdraw_native_preview

    job = job_at_reconstruct["job"]
    assert job_at_reconstruct["reconstruct"]()["job_state"] == "cp3_native_preview"
    finding = "S01: a white box sits behind the title in the render"
    withdraw_native_preview(job, [finding], reason="visual review")
    write_json_atomic(service_path(job, "reconstruct-issues.json"), {"issues": [finding, "the previous session failed: x"]})
    result = job_at_reconstruct["reconstruct"]()
    assert result["job_state"] == "cp3_native_preview"
    attempts = read_json(service_path(job, "reconstruct-report.json"))["attempts"]
    assert attempts[0]["attempt"] == 1  # no Worker-only acceptance
    prompts = sorted(job_at_reconstruct["home"].glob("prompt-*.txt"))
    assert finding in prompts[-1].read_text(encoding="utf-8")


def test_visible_text_ocr_misreads_is_confirmed_by_ink_and_covered_text_is_not(tmp_path):
    # Measured 2026-10-18: OCR read a visible white-on-dark title as components (宀 匕 石 开 京); the covered
    # r3 S07 title kept 5 % of the master's ink, visible native lines 76-107 %.
    from PIL import ImageDraw
    from autoppt_service.reconstruct import render_visibility_errors

    def page(path, *, text, shift=0):
        image = Image.new("RGB", (400, 200), (24, 64, 96))
        if text:
            draw = ImageDraw.Draw(image)
            for k in range(6):
                draw.rectangle((40 + shift + 22 * k, 40, 52 + shift + 22 * k, 70), fill=(255, 255, 255))
        image.save(path)

    page(tmp_path / "master.png", text=True)
    page(tmp_path / "visible.png", text=True, shift=2)
    page(tmp_path / "covered.png", text=False)

    def ocr(image, out_json):
        if Path(image).name == "master.png":
            return {"width": 400, "height": 200,
                    "lines": [{"text": "研 究 背 景", "bbox": {"x": 40, "y": 40, "w": 130, "h": 30}}]}
        return {"lines": [{"text": "宀 匕 石 开 京"}]}

    targets = {"S01": ["研究背景"]}
    masters = {"S01": tmp_path / "master.png"}
    assert render_visibility_errors({"S01": tmp_path / "visible.png"}, targets, ocr=ocr, masters=masters,
                                    work_dir=tmp_path / "a") == []
    errors = render_visibility_errors({"S01": tmp_path / "covered.png"}, targets, ocr=ocr, masters=masters,
                                      work_dir=tmp_path / "b")
    assert errors and "研究背景" in errors[0]
