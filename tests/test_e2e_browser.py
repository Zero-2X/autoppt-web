"""Browser end-to-end: a live API, a mock worker thread and the real UI in headless Chromium.

Covers login, redeeming credits, the order wizard with a real file upload, all three
checkpoints (including a paid redo confirmed in a dialog), delivery and a download.
Skipped when Playwright or its Chromium is not installed.
"""

from __future__ import annotations

import io
import socket
import threading
import time
from pathlib import Path

import httpx
import pytest

playwright_api = pytest.importorskip("playwright.sync_api")


def free_port() -> int:
    with socket.socket() as sock:
        sock.bind(("127.0.0.1", 0))
        return sock.getsockname()[1]


@pytest.fixture()
def live(tmp_path):
    import uvicorn

    from autoppt_api.app import create_app
    from autoppt_api.config import Settings
    from autoppt_service.worker_client import ApiWorker

    port = free_port()
    settings = Settings(database_url=f"sqlite:///{(tmp_path / 'api.db').as_posix()}", storage_dir=tmp_path / "storage",
                        dev_echo_codes=True, admin_emails=("admin@example.com",))
    server = uvicorn.Server(uvicorn.Config(create_app(settings), host="127.0.0.1", port=port, log_level="warning"))
    thread = threading.Thread(target=server.run, daemon=True)
    thread.start()
    base = f"http://127.0.0.1:{port}"
    for _ in range(100):
        try:
            if httpx.get(f"{base}/api/health", timeout=1).status_code == 200:
                break
        except httpx.HTTPError:
            time.sleep(0.1)
    http = httpx.Client(base_url=base, timeout=60)
    code = http.post("/api/auth/code", json={"email": "admin@example.com"}).json()["dev_code"]
    admin = {"Authorization": "Bearer " + http.post("/api/auth/login", json={"email": "admin@example.com", "code": code}).json()["access_token"]}
    token = http.post("/api/admin/workers", json={"name": "e2e", "capabilities": ["all"]}, headers=admin).json()["token"]
    redeem = http.post("/api/admin/redeem-codes", json={"credits": 400, "note": "e2e"}, headers=admin).json()["codes"][0]
    stop = threading.Event()

    def work() -> None:
        with httpx.Client(base_url=base, timeout=60) as client:
            worker = ApiWorker(client, token, tmp_path / "worker", mock=True, heartbeat_s=3600)
            while not stop.is_set():
                if worker.run_once() is None:
                    time.sleep(0.5)

    worker_thread = threading.Thread(target=work, daemon=True)
    worker_thread.start()
    yield {"base": base, "redeem": redeem, "tmp": tmp_path}
    stop.set()
    worker_thread.join(timeout=30)
    server.should_exit = True
    thread.join(timeout=10)


def thesis_file(tmp_path: Path) -> Path:
    from docx import Document

    document = Document()
    document.add_heading("研究背景", level=1)
    document.add_paragraph("城市内涝频发，现有预警只能给出区县级结果，平均提前 18 分钟。")
    path = tmp_path / "论文终稿.docx"
    document.save(path)
    return path


def test_order_review_and_download_in_a_real_browser(live):
    with playwright_api.sync_playwright() as p:
        try:
            browser = p.chromium.launch()
        except Exception as exc:  # noqa: BLE001 - browser binaries missing
            pytest.skip(f"chromium not available: {exc}")
        page = browser.new_page(accept_downloads=True)
        dialogs: list[str] = []
        page.on("dialog", lambda dialog: (dialogs.append(dialog.message), dialog.accept()))
        page.set_default_timeout(60_000)

        # Sign in with an email code (the dev server shows the code on the page).
        page.goto(f"{live['base']}/#/login")
        page.get_by_label("邮箱").fill("student@example.com")
        page.get_by_role("button", name="发送验证码").click()
        code = page.get_by_text("开发模式验证码").inner_text().split("：")[-1].strip("）")
        page.get_by_label("验证码").fill(code)
        page.get_by_role("button", name="登录", exact=True).click()
        page.wait_for_url("**/#/jobs")

        # Credits from a redeem code.
        page.goto(f"{live['base']}/#/credits")
        page.get_by_label("兑换码").fill(live["redeem"])
        page.get_by_role("button", name="兑换").click()
        page.get_by_text("已到账 400 积分").wait_for()

        # Order wizard with a real file upload.
        page.goto(f"{live['base']}/#/new")
        page.get_by_role("button", name="下一步", exact=True).click()
        page.locator("input[type=file]").set_input_files(str(thesis_file(live["tmp"])))
        page.get_by_text("已上传（docx）").wait_for()
        page.get_by_label("描述文字").fill("硕士论文答辩，重点讲研究问题、方法和本人贡献。")
        page.get_by_role("button", name="下一步", exact=True).click()
        page.get_by_role("button", name="下一步", exact=True).click()
        page.get_by_label("标题", exact=True).fill("城市内涝预警研究")
        page.get_by_text("接受用 Office 字体").click()
        page.get_by_role("button", name="下一步", exact=True).click()
        submit = page.get_by_role("button", name="确认下单")
        submit.wait_for()
        assert "244" in submit.inner_text()
        submit.click()
        page.wait_for_url("**/#/jobs/*")

        # CP1: approve the outline.
        page.get_by_text("大纲确认（第 1 轮）").wait_for()
        page.get_by_role("button", name="通过", exact=True).click()

        # CP2: redo S03 (a paid redo confirmed in a dialog), then approve round 2.
        page.get_by_text("母版审核（第 1 轮）").wait_for()
        page.get_by_label("S03 重做").check()
        page.get_by_placeholder("重做原因（勾选的页）").fill("S03 层次不清楚")
        page.get_by_role("button", name="重做所选页").click()
        page.get_by_text("母版审核（第 2 轮）").wait_for()
        assert any("扣除 4 积分" in message for message in dialogs)
        page.get_by_role("button", name="通过", exact=True).click()

        # CP3: approval needs the user's own sentence.
        page.get_by_text("原生预览批准（第 1 轮）").wait_for()
        page.get_by_role("button", name="接受并交付").click()
        assert any("写一句话" in message for message in dialogs)
        page.get_by_placeholder("请用一句话写下你的决定").fill("接受预览中的非文字差异，优先真实可编辑")
        page.get_by_role("button", name="接受并交付").click()

        # Delivered: download the native deck.
        page.get_by_text("已交付").first.wait_for()
        with page.expect_download() as download_info:
            page.get_by_role("link", name="native-deck.pptx").click()
        downloaded = Path(download_info.value.path())
        from pptx import Presentation

        assert len(Presentation(io.BytesIO(downloaded.read_bytes())).slides) > 0

        # Credits after delivery: nothing left frozen.
        page.goto(f"{live['base']}/#/credits")
        page.get_by_text("积分流水").wait_for()
        text = page.locator("main").inner_text()
        assert "冻结\n0" in text
        browser.close()
