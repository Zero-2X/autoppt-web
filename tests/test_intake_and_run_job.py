"""intake.py, previews.py and the run_job --mock flow."""

from __future__ import annotations

import json
import zipfile
from pathlib import Path

import pytest
from PIL import Image

from autoppt_service.checkpoint import pending_checkpoint, submit_response
from autoppt_service.intake import IntakeError, run_intake
from autoppt_service.jobfiles import read_json, service_path
from autoppt_service.previews import build_master_previews, build_native_previews, checkpoint_previews, diff_metrics
from autoppt_service.repo_paths import SERVICE_ROOT
from autoppt_service.run_job import PhaseError, PhaseNotReady, main as run_job_main, run_all, run_phase
from autoppt_service.status import load_job_state
from autoppt_workflow.presentation.schemas import normalize_evidence_ledger

DESCRIPTION = "本研究面向城市内涝预警，目标是在暴雨前两小时给出街区级风险等级，并说明本人负责的模型部分。" * 6


def make_request(tier="native_editable", files=None, **extra) -> dict:
    request = {
        "schema_version": "autoppt-job-request-v1",
        "service": "new_deck",
        "inputs": {"files": files or [], "description": DESCRIPTION, "must_include": ["本人贡献单独一页"]},
        "scenario": "thesis_defense",
        "slides": {"mode": "fixed", "count": 3},
        "cover": {"title": "城市内涝预警研究", "author": "张三"},
        "delivery": {"tier": tier, "font_policy": "office_native_yahei"} if tier == "native_editable" else {"tier": tier},
    }
    request.update(extra)
    return request


@pytest.fixture()
def uploads(tmp_path):
    from docx import Document

    src = tmp_path / "uploads"
    src.mkdir()
    document = Document()
    document.add_heading("研究背景", level=1)
    document.add_paragraph("城市内涝频发，现有预警只能给出区县级结果。")
    document.save(src / "paper.docx")
    Image.new("RGB", (40, 30), "green").save(src / "fig.png")
    files = [
        {"file_id": "f_01", "role": "primary_source", "name": "论文终稿.docx"},
        {"file_id": "f_02", "role": "keep_as_image", "name": "图3-2.png"},
    ]
    return files, {"f_01": str(src / "paper.docx"), "f_02": str(src / "fig.png")}


# ---------------------------------------------------------------- intake


def test_intake_builds_job_directory(tmp_path, uploads):
    files, paths = uploads
    job = tmp_path / "job"
    report = run_intake(make_request(files=files), paths, job, job_id="J1")
    assert report["status"] in ("ready", "ready_with_warnings"), report
    assert report["job_state"] == "planning"
    assert sorted(p.name for p in (job / "source-materials").iterdir()) == ["f_01__论文终稿.docx", "f_02__图3-2.png"]

    ledger = read_json(job / "workspace/evidence_workspace/evidence-ledger.json")
    normalized = normalize_evidence_ledger([ledger])  # the workflow's own reader
    assert normalized["summary"]["count"] == len(ledger["evidence_items"]) >= 3
    manifest = read_json(job / "workspace/evidence_workspace/source-manifest.json")
    assert all(not Path(record["path"]).is_absolute() for record in manifest["files"])
    assert all(not Path(asset["path"]).is_absolute() for asset in manifest["assets"])

    brief = read_json(job / "workspace/presentation_handoff/planning-brief.json")
    assert brief["scenario"]["profile_id"] == "thesis_defense"
    assert brief["slides"]["target"] == 3
    assert {"field": "title", "text": "城市内涝预警研究"} in brief["exact_text_whitelist"]
    assert brief["keep_as_image"][0]["source_file"] == "图3-2.png"
    assert brief["checkpoints"] == ["outline", "master", "native_preview"]
    assert "城市内涝预警研究" in (job / "workspace/presentation_handoff/planning-brief.md").read_text(encoding="utf-8")
    assert read_json(service_path(job, "job-request.json"))["style"]["profile"] == "academic_light"


def test_intake_rejects_invalid_request(tmp_path):
    bad = make_request()
    bad["delivery"]["font_policy"] = "master_glyphs"
    report = run_intake(bad, {}, tmp_path / "job", job_id="J1")
    assert report["status"] == "rejected"
    assert {e["code"] for e in report["errors"]} == {"office_font_policy_required"}
    assert load_job_state(tmp_path / "job") is None


def test_intake_blocks_on_missing_or_rejected_upload_and_reruns(tmp_path, uploads):
    files, paths = uploads
    job = tmp_path / "job"
    broken = tmp_path / "broken.docx"
    broken.write_bytes(b"not a docx")
    report = run_intake(make_request(files=files), {"f_01": str(broken), "f_02": paths["f_02"]}, job, job_id="J1")
    assert report["status"] == "blocked"
    assert report["rejected_files"][0]["file_id"] == "f_01"
    state = load_job_state(job)
    assert state["state"] == "blocked" and state["blocked_from"] == "intake"

    report = run_intake(make_request(files=files), paths, job, job_id="J1")  # user replaced the upload
    assert report["job_state"] == "planning"
    assert not any("broken" in p.name for p in (job / "source-materials").iterdir())
    with pytest.raises(IntakeError):
        run_intake(make_request(files=files), paths, job, job_id="J1")  # already past intake


def test_intake_refuses_job_dir_inside_a_repository(tmp_path):
    with pytest.raises(IntakeError):
        run_intake(make_request(), {}, SERVICE_ROOT / "jobs" / "J1", job_id="J1")
    with pytest.raises(IntakeError):
        run_intake(make_request(), {}, tmp_path / "repo" / "job", job_id="J1", protected_roots=[tmp_path / "repo"])


# ---------------------------------------------------------------- previews


def slide(path: Path, colour, size=(160, 90)) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", size, colour).save(path)
    return path


def test_previews(tmp_path):
    masters = {sid: slide(tmp_path / f"assets/slides/{sid}.png", "white") for sid in ("S01", "S02")}
    natives = {"S01": slide(tmp_path / "render/S01.png", "white", (320, 180)), "S02": slide(tmp_path / "render/S02.png", "black")}
    cp2 = build_master_previews(tmp_path, masters)
    kinds = [entry["kind"] for entry in cp2["previews"]]
    assert kinds.count("thumbnail") == 2 and kinds.count("contact_sheet") == 1 and kinds.count("slide") == 2
    cp3 = build_native_previews(tmp_path, masters, natives)
    metrics = cp3["diff_metrics"]
    assert metrics["S01"]["changed_pixel_ratio"] == 0  # same content at a different size
    assert metrics["S02"]["changed_pixel_ratio"] == 1
    assert (tmp_path / "service/previews/native/manifest.json").exists()
    for entry in checkpoint_previews(cp3):
        assert set(entry) <= {"path", "kind", "slide_id", "sha256"}
        assert (tmp_path / entry["path"]).is_file()
    with pytest.raises(ValueError):
        build_native_previews(tmp_path, masters, {"S01": natives["S01"]})


def test_diff_metrics_partial_change():
    a = Image.new("RGB", (10, 10), "white")
    b = a.copy()
    b.paste((0, 0, 0), (0, 0, 5, 10))
    _, metrics = diff_metrics(a, b)
    assert metrics["changed_pixel_ratio"] == pytest.approx(0.5)


# ---------------------------------------------------------------- run_job


def test_mock_flow_native_tier_end_to_end(tmp_path, uploads):
    files, paths = uploads
    job = tmp_path / "job"
    results = run_all(job, job_id="J1", request=make_request(files=files), uploads=paths, auto_respond=True)
    assert [r["phase"] for r in results] == ["intake", "planning", "masters", "reconstruct", "package"]
    state = load_job_state(job)
    assert state["state"] == "delivered"
    assert [h["to"] for h in state["history"]] == [
        "queued", "intake", "planning", "cp1_outline", "masters", "cp2_masters", "image_deck_ready",
        "reconstructing", "office_qa", "cp3_native_preview", "packaging", "delivered",
    ]
    acceptance = read_json(service_path(job, "acceptance.json"))
    assert acceptance["mock"] is True and acceptance["tier"] != "native-priority-approved"
    manifest = read_json(job / "output/delivery-manifest.json")
    assert manifest["mock"] is True and "MOCK" in manifest["notice"]
    with zipfile.ZipFile(job / "output/J1-MOCK.zip") as bundle:
        assert "output/native-deck.pptx" in bundle.namelist()
    usage = read_json(service_path(job, "usage-summary.json"))
    assert set(usage["phases"]) == {"intake", "planning", "masters", "reconstruct", "package"}
    assert len(list((service_path(job, "checkpoints")).iterdir())) == 3


def test_mock_flow_stops_at_checkpoints_and_handles_rework(tmp_path, uploads):
    files, paths = uploads
    job = tmp_path / "job"
    run_all(job, job_id="J1", request=make_request("image_deck", files=files), uploads=paths, auto_respond=False)
    assert load_job_state(job)["state"] == "cp1_outline"
    with pytest.raises(PhaseError, match="waiting"):
        run_phase(job, "masters", job_id="J1", mock=True)

    cp1 = pending_checkpoint(job)
    submit_response(job, {"checkpoint_id": cp1["checkpoint_id"], "decision": "revise", "responder": "user",
                          "user_reply": "第二页改成本人贡献", "annotations": [{"slide_id": "S02", "text": "改成本人贡献"}]})
    run_phase(job, "planning", job_id="J1", mock=True)  # back to planning, round 2
    cp1 = pending_checkpoint(job)
    assert cp1["round"] == 2
    brief = read_json(job / "workspace/presentation_handoff/slide_brief.json")
    assert brief["slides"][1]["revision_note"] == "改成本人贡献"
    submit_response(job, {"checkpoint_id": cp1["checkpoint_id"], "decision": "approve", "responder": "user", "user_reply": "可以"})

    run_phase(job, "masters", job_id="J1", mock=True)
    cp2 = pending_checkpoint(job)
    before = {sid: digest for sid, digest in cp2["master_sha256"].items()}
    submit_response(job, {"checkpoint_id": cp2["checkpoint_id"], "decision": "redo_slides", "responder": "user",
                          "user_reply": "S01 重做", "redo_slide_ids": ["S01"]})
    run_phase(job, "masters", job_id="J1", mock=True)
    cp2 = pending_checkpoint(job)
    assert cp2["round"] == 2 and cp2["master_sha256"]["S02"] == before["S02"]
    submit_response(job, {"checkpoint_id": cp2["checkpoint_id"], "decision": "approve", "responder": "user", "user_reply": "通过"})

    result = run_phase(job, "package", job_id="J1", mock=True)
    assert result["job_state"] == "delivered"
    assert (job / "output/image-deck.pptx").exists()


def test_real_routes_report_what_they_need(tmp_path, uploads, monkeypatch):
    files, paths = uploads
    job = tmp_path / "job"
    monkeypatch.delenv("AUTOPPT_WORKER_TEMPLATE", raising=False)
    run_phase(job, "intake", job_id="J1", request=make_request(files=files), uploads=paths)
    with pytest.raises(PhaseNotReady, match="worker template"):
        run_phase(job, "planning", job_id="J1")
    with pytest.raises(PhaseError):
        run_phase(job, "reconstruct", job_id="J1", mock=True)  # wrong state


def test_cli_respond_and_all(tmp_path, uploads, capsysbinary):
    files, paths = uploads
    request_path = tmp_path / "request.json"
    request_path.write_text(json.dumps(make_request("planning_package", files=files), ensure_ascii=False), encoding="utf-8")
    uploads_path = tmp_path / "uploads.json"
    uploads_path.write_text(json.dumps(paths), encoding="utf-8")
    job = tmp_path / "job"
    common = ["--job-dir", str(job), "--job-id", "J1"]
    assert run_job_main([str(request_path), *common, "--phase", "all", "--mock", "--uploads", str(uploads_path)]) == 0
    capsysbinary.readouterr()
    cp1 = pending_checkpoint(job)
    answer = tmp_path / "answer.json"
    answer.write_text(json.dumps({"checkpoint_id": cp1["checkpoint_id"], "decision": "approve", "responder": "user", "user_reply": "好"}, ensure_ascii=False), encoding="utf-8")
    assert run_job_main([*common, "--respond", str(answer)]) == 0
    assert run_job_main([*common, "--phase", "package", "--mock"]) == 0
    assert load_job_state(job)["state"] == "delivered"
    with pytest.raises(PhaseError):
        run_job_main([*common, "--phase", "planning"])  # already delivered


def test_paid_addons_are_delivered(tmp_path, uploads):
    from pptx import Presentation

    files, paths = uploads
    request = make_request(files=files)
    request["delivery"].update(speaker_notes=True, anticipated_questions=True)
    job = tmp_path / "job"
    run_all(job, job_id="J1", request=request, uploads=paths, auto_respond=True)
    assert load_job_state(job)["state"] == "delivered"
    manifest = read_json(job / "output/delivery-manifest.json")
    paths_out = {item["path"] for item in manifest["files"]}
    assert {"output/speaker-notes.md", "output/anticipated-questions.md"} <= paths_out
    questions = (job / "output/anticipated-questions.md").read_text(encoding="utf-8")
    assert questions.count("## ") == 8
    deck = Presentation(str(job / "output/native-deck.pptx"))
    assert deck.slides[0].notes_slide.notes_text_frame.text.startswith("[MOCK] S01")


def test_planning_rejects_a_plan_without_ordered_addons():
    from autoppt_service.addons import addon_errors

    plan = {"slides": [{"slide_id": "S01", "speaker_notes": "讲稿"}, {"slide_id": "S02"}], "anticipated_questions": [{"question": "q"}]}
    errors = addon_errors(plan, ["speaker_notes", "anticipated_questions"])
    assert "S02" in errors[0] and "at least 8" in errors[1]
    assert addon_errors(plan, []) == []
