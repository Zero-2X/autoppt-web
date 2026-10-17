"""Real planning phase with a fake ``codex`` CLI that plays the planning agent."""

from __future__ import annotations

import json
import sys
from pathlib import Path

import pytest

from autoppt_service.checkpoint import pending_checkpoint, submit_response
from autoppt_service.jobfiles import read_json, read_jsonl, service_path
from autoppt_service.planning import EVIDENCE_LABELS, validate_deck_plan
from autoppt_service.run_job import PhaseNotReady, run_phase
from autoppt_service.status import load_job_state
from autoppt_service.worker_template import build_template

# The fake agent reads the brief and ledger from its cwd (the job directory),
# writes deck-plan.json and status.json, and reports usage like codex exec.
FAKE_PLANNER = r'''
import json, os, sys, uuid
from pathlib import Path

prompt = sys.stdin.buffer.read().decode("utf-8")
home = Path(os.environ["CODEX_HOME"])
counter = home / "calls.txt"
calls = int(counter.read_text()) + 1 if counter.exists() else 1
counter.write_text(str(calls))
(home / f"prompt-{calls}.txt").write_text(prompt, encoding="utf-8")
(home / f"argv-{calls}.json").write_text(json.dumps(sys.argv[1:]), encoding="utf-8")
mode = os.environ.get("FAKE_PLAN_MODE", "ok")
print(json.dumps({"type": "thread.started", "thread_id": str(uuid.uuid4())}))
if mode == "rate_limit_once" and calls == 1:
    print(json.dumps({"type": "error", "message": "You've hit your usage limit."}))
    sys.exit(1)
job = Path.cwd()
assert (job / "AGENTS.md").exists(), "AGENTS.md must be rendered before the session"
def status(state, **extra):
    payload = {"schema_version": "autoppt-agent-status-v1", "phase": "planning", "state": state, "updated_at": "2027-01-01T00:00:00Z", **extra}
    (job / "service" / "status.json").write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
if mode == "blocked":
    status("blocked", blocked_reason="主资料只有目录页，缺少正文")
    print(json.dumps({"type": "turn.completed", "usage": {"input_tokens": 10, "output_tokens": 1}}))
    sys.exit(0)
brief = json.loads((job / "workspace/presentation_handoff/planning-brief.json").read_text(encoding="utf-8"))
ledger = json.loads((job / "workspace/evidence_workspace/evidence-ledger.json").read_text(encoding="utf-8"))
evidence = [i["evidence_id"] for i in ledger["evidence_items"] if i["item_type"] != "user_statement"]
count = brief["slides"]["target"]
cover = [entry["text"] for entry in brief["exact_text_whitelist"]]
must = brief["must_include"]
slides = []
for index in range(1, count + 1):
    text = cover if index == 1 else [f"第{index}页结论"]
    if index == 2:
        text = text + must
        if mode == "bad_then_ok" and "does not appear in any evidence item" not in prompt:
            text = text + ["准确率提升到 97.3%"]
    slides.append({
        "slide_id": f"S{index:02d}", "role": brief["scenario"]["page_roles"][min(index - 1, 3)],
        "headline": text[0], "message": f"第{index}页回答的问题", "exact_text": text,
        "visual_brief": "左侧证据图，右侧结论", "evidence_summary": "来自论文第一节",
        "source_refs": [evidence[0]] if evidence else ["framing"],
        "source_assets": [a["asset_id"] for a in brief["keep_as_image"]] if index == 3 else [],
        "speaker_notes": ("REVISED: " + prompt.split("<<<USER_REPLY\n", 1)[1].split("\nUSER_REPLY>>>", 1)[0]) if "<<<USER_REPLY" in prompt and index == 2 else "",
    })
plan = {
    "schema_version": "autoppt-deck-plan-v1",
    "deck": {"title": cover[0], "use_case": "thesis defense"},
    "slides": slides,
    "must_include_coverage": [{"item": item, "slide_ids": ["S02"]} for item in must],
    "learning_context": {"files_read": ["improvement/ppt-improvement-ledger.jsonl"], "lessons_applied": ["one conclusion per slide"]},
}
(job / "workspace/presentation_handoff/deck-plan.json").write_text(json.dumps(plan, ensure_ascii=False), encoding="utf-8")
status("done", progress=1.0)
print(json.dumps({"type": "turn.completed", "usage": {"input_tokens": 1200, "cached_input_tokens": 200, "output_tokens": 300}}))
'''

DESCRIPTION = "本研究面向城市内涝预警，目标是在暴雨前两小时给出街区级风险等级，并说明本人负责的模型部分。" * 6


@pytest.fixture()
def setup(tmp_path, monkeypatch):
    from docx import Document
    from PIL import Image
    from test_worker_template import fake_repo

    template = build_template(tmp_path / "templates", repo=fake_repo(tmp_path / "autoppt"), version="v1")
    script = tmp_path / "fake_codex.py"
    script.write_text(FAKE_PLANNER, encoding="utf-8")
    home = tmp_path / "homes" / "acct1"
    home.mkdir(parents=True)
    profiles = tmp_path / "profiles.json"
    profiles.write_text(json.dumps([{"name": "acct1", "codex_home": str(home)}, {"name": "acct2", "codex_home": str(home)}]), encoding="utf-8")

    uploads = tmp_path / "uploads"
    uploads.mkdir()
    document = Document()
    document.add_heading("研究背景", level=1)
    document.add_paragraph("城市内涝频发，现有预警只能给出区县级结果，平均提前 30 分钟。")
    document.save(uploads / "paper.docx")
    Image.new("RGB", (40, 30), "green").save(uploads / "fig.png")
    request = {
        "schema_version": "autoppt-job-request-v1",
        "service": "new_deck",
        "inputs": {
            "files": [
                {"file_id": "f_01", "role": "primary_source", "name": "论文终稿.docx"},
                {"file_id": "f_02", "role": "keep_as_image", "name": "图3-2.png"},
            ],
            "description": DESCRIPTION,
            "must_include": ["本人贡献单独一页"],
            "must_avoid": ["把伪标签一致性写成实测精度"],
        },
        "scenario": "thesis_defense",
        "slides": {"mode": "fixed", "count": 4},
        "cover": {"title": "城市内涝预警研究", "author": "张三", "date": "2027-05"},
        "delivery": {"tier": "native_editable", "font_policy": "office_native_yahei"},
    }
    job = tmp_path / "job"
    run_phase(job, "intake", job_id="J1", request=request,
              uploads={"f_01": str(uploads / "paper.docx"), "f_02": str(uploads / "fig.png")})
    assert load_job_state(job)["state"] == "planning"

    def plan(mode="ok", **extra):
        monkeypatch.setenv("FAKE_PLAN_MODE", mode)
        return run_phase(job, "planning", job_id="J1", template=template, profiles_path=profiles,
                         command=[sys.executable, str(script)], **extra)

    return {"job": job, "home": home, "plan": plan, "template": template, "request": request}


def test_real_planning_opens_cp1_and_builds_workflow_inputs(setup):
    job = setup["job"]
    result = setup["plan"]()
    assert result["job_state"] == "cp1_outline", result
    checkpoint = pending_checkpoint(job)
    assert checkpoint["type"] == "outline" and checkpoint["round"] == 1
    assert checkpoint["previews"][0]["path"] == "workspace/presentation_handoff/outline-preview.md"

    prompts = read_json(job / "image-prompts.json")
    assert prompts["prompt_policy"]["builtin_imagegen_only"] is True
    assert prompts["style_profile"] == "academic_light"  # fixed by the Worker, not the agent
    assert [s["final_path"] for s in prompts["slides"]] == [f"assets/slides/S{i:02d}.png" for i in range(1, 5)]
    assert "城市内涝预警研究" in prompts["slides"][0]["prompt_zh"]
    assert "Microsoft YaHei" in prompts["slides"][0]["prompt_zh"]
    spec = read_json(job / "deck-spec.json")
    assert spec["deck"]["output_mode"] == "imagegen_full_slide" and len(spec["slides"]) == 4
    brief = read_json(job / "workspace/presentation_handoff/slide_brief.json")
    assert brief["slides"][1]["title"] == "第2页结论"

    agents = (job / "AGENTS.md").read_text(encoding="utf-8")
    assert "Job `J1`" in agents and setup["template"].resolve().as_posix() in agents
    turns = [r for r in read_jsonl(service_path(job, "usage.jsonl")) if r.get("kind") == "agent_turn"]
    assert turns[-1]["phase"] == "planning" and turns[-1]["usage"]["input_tokens"] == 1200
    report = read_json(service_path(job, "planning-report.json"))
    assert report["learning_context"]["files_read"]
    argv = json.loads((setup["home"] / "argv-1.json").read_text(encoding="utf-8"))
    assert argv[argv.index("-s") + 1] == "workspace-write"  # writes stay inside the job directory
    assert "--search" not in argv and "--ephemeral" not in argv


def test_fabricated_number_is_sent_back_then_fixed(setup):
    result = setup["plan"]("bad_then_ok")
    assert result["job_state"] == "cp1_outline"
    report = read_json(service_path(setup["job"], "planning-report.json"))
    assert len(report["attempts"]) == 2
    assert any("97.3" in error for error in report["attempts"][0]["errors"])
    second_prompt = (setup["home"] / "prompt-2.txt").read_text(encoding="utf-8")
    assert "does not appear in any evidence item" in second_prompt


def test_agent_blocked_moves_job_to_blocked(setup):
    result = setup["plan"]("blocked")
    assert result["job_state"] == "blocked"
    state = load_job_state(setup["job"])
    assert "缺少正文" in state["blocked_reason"] and state["blocked_from"] == "planning"


def test_rate_limited_profile_rests_and_next_profile_runs(setup):
    result = setup["plan"]("rate_limit_once")
    assert result["job_state"] == "cp1_outline"
    attempts = read_json(service_path(setup["job"], "planning-report.json"))["attempts"]
    assert [a["profile"] for a in attempts] == ["acct1", "acct2"]


def test_revision_round_passes_user_words(setup):
    job = setup["job"]
    setup["plan"]()
    cp1 = pending_checkpoint(job)
    submit_response(job, {"checkpoint_id": cp1["checkpoint_id"], "decision": "revise", "responder": "user",
                          "user_reply": "第二页加上实验对比"})
    result = setup["plan"]()
    assert result["job_state"] == "cp1_outline"
    assert pending_checkpoint(job)["round"] == 2
    plan = read_json(job / "workspace/presentation_handoff/deck-plan.json")
    assert plan["slides"][1]["speaker_notes"] == "REVISED: 第二页加上实验对比"
    assert (job / "service/codex-runs/planning-r2-a1").is_dir()


def test_missing_template_or_4_3_is_reported(setup, monkeypatch):
    monkeypatch.delenv("AUTOPPT_WORKER_TEMPLATE", raising=False)
    with pytest.raises(PhaseNotReady, match="worker template"):
        run_phase(setup["job"], "planning", job_id="J1")
    request = read_json(service_path(setup["job"], "job-request.json"))
    request["delivery"]["aspect_ratio"] = "4:3"
    service_path(setup["job"], "job-request.json").write_text(json.dumps(request, ensure_ascii=False), encoding="utf-8")
    result = setup["plan"]()
    assert result["job_state"] == "blocked" and "16:9" in result["issues"][0]


# ---------------------------------------------------------------- validator


def good_plan() -> tuple[dict, dict, dict, dict]:
    brief = {
        "slides": {"min": 2, "max": 3, "target": 2},
        "exact_text_whitelist": [{"field": "title", "text": "城市内涝预警研究"}, {"field": "date", "text": "2027-05"}],
        "must_include": ["本人贡献单独一页"],
        "must_avoid": ["实测精度"],
        "keep_as_image": [{"asset_id": "f_02_image", "source_file": "fig.png"}],
    }
    ledger = {"evidence_items": [{"evidence_id": "ev_f_01_sec001", "fact_or_data": "预警平均提前 30 分钟，覆盖 1,200 个街区"}]}
    manifest = {"assets": [{"asset_id": "f_02_image"}]}
    plan = {
        "schema_version": "autoppt-deck-plan-v1",
        "deck": {"title": "城市内涝预警研究"},
        "slides": [
            {"slide_id": "S01", "role": "cover", "headline": "城市内涝预警研究", "message": "m",
             "exact_text": ["城市内涝预警研究", "2027-05"], "visual_brief": "v", "evidence_summary": "e", "source_refs": ["framing"]},
            {"slide_id": "S02", "role": "result", "headline": "提前 30 分钟", "message": "m",
             "exact_text": ["提前 30 分钟覆盖 1200 个街区", "本人贡献单独一页", "第 1 步"], "visual_brief": "v",
             "evidence_summary": "e", "source_refs": ["ev_f_01_sec001"], "source_assets": ["f_02_image"]},
        ],
        "must_include_coverage": [{"item": "本人贡献单独一页", "slide_ids": ["S02"]}],
        "learning_context": {"files_read": ["ledger"]},
    }
    return plan, brief, ledger, manifest


def test_validator_accepts_good_plan():
    errors, warnings = validate_deck_plan(*good_plan())
    assert errors == [] and warnings == []
    assert "framing" in EVIDENCE_LABELS


@pytest.mark.parametrize(
    "mutate, expected",
    [
        (lambda p, b: p["slides"][1]["source_refs"].append("ev_missing"), "unknown evidence ids"),
        (lambda p, b: p["slides"][0]["exact_text"].remove("2027-05"), "2027-05"),
        (lambda p, b: p["slides"][1]["exact_text"].append("实测精度 很高"), "must-avoid"),
        (lambda p, b: p.update(must_include_coverage=[]), "not mapped"),
        (lambda p, b: p["slides"][1]["exact_text"].append("提升 45%"), "number 45"),
        (lambda p, b: p["slides"].append(dict(p["slides"][1], slide_id="S03")) or p["slides"].append(dict(p["slides"][1], slide_id="S04")), "outside the range"),
        (lambda p, b: p["slides"][1].update(slide_id="S05"), "in order"),
        (lambda p, b: p["slides"][1].update(source_assets=["nope"]), "unknown source_assets"),
        (lambda p, b: p.pop("learning_context"), "learning_context"),
    ],
)
def test_validator_rejects(mutate, expected):
    plan, brief, ledger, manifest = good_plan()
    mutate(plan, brief)
    errors, _ = validate_deck_plan(plan, brief, ledger, manifest)
    assert any(expected in error for error in errors), errors


def test_validator_warnings():
    plan, brief, ledger, manifest = good_plan()
    plan["slides"][1]["source_assets"] = []
    plan["slides"].extend([dict(plan["slides"][1], slide_id="S03"), dict(plan["slides"][1], slide_id="S04")])
    plan["outline_conflicts"] = ["结果部分需要四页才能讲完"]
    plan["slides"][0]["exact_text"].append("覆盖 1200 个街区")  # in the ledger, not cited on S01
    errors, warnings = validate_deck_plan(plan, brief, ledger, manifest)
    assert errors == []
    assert any("outside" in w for w in warnings)
    assert any("not used" in w for w in warnings)
    assert any("not in this slide's cited items" in w for w in warnings)


def test_usage_limit_on_every_profile_defers_instead_of_blocking(setup, monkeypatch, tmp_path):
    # One profile only, and the fake CLI reports a usage limit on its first call.
    profiles = tmp_path / "one-profile.json"
    profiles.write_text(json.dumps([{"name": "acct1", "codex_home": str(setup["home"])}]), encoding="utf-8")
    monkeypatch.setenv("FAKE_PLAN_MODE", "rate_limit_once")
    script = tmp_path / "fake_codex.py"
    result = run_phase(setup["job"], "planning", job_id="J1", template=setup["template"], profiles_path=profiles,
                       command=[sys.executable, str(script)])
    assert result["deferred"] is True and result["job_state"] == "planning"
    assert read_json(service_path(setup["job"], "deferred.json"))["phase"] == "planning"
    assert load_job_state(setup["job"])["state"] == "planning"  # not blocked: no user action needed
    result = setup["plan"]()  # after the reset the same job continues
    assert result["job_state"] == "cp1_outline"
    assert not service_path(setup["job"], "deferred.json").exists()


def test_retry_hint_and_rollout_usage(tmp_path):
    from autoppt_service.planning import usage_limit_retry_hint
    from autoppt_service.usage import rollout_token_usage

    message = "You've hit your usage limit. Upgrade to Pro, visit the usage page or try again at 2:21 PM."
    assert usage_limit_retry_hint([message]) == "2:21 PM"
    assert usage_limit_retry_hint(["boom"]) is None
    rollout = tmp_path / "rollout.jsonl"
    rollout.write_text(
        "\n".join(
            json.dumps(line)
            for line in [
                {"type": "event_msg", "payload": {"type": "token_count", "info": {"total_token_usage": {"input_tokens": 10}}}},
                {"type": "event_msg", "payload": {"type": "token_count", "info": None}},
                {"type": "event_msg", "payload": {"type": "token_count", "info": {"total_token_usage": {"input_tokens": 99, "output_tokens": 3}}}},
            ]
        ),
        encoding="utf-8",
    )
    assert rollout_token_usage(rollout) == {"input_tokens": 99, "output_tokens": 3}
    assert rollout_token_usage(None) == {}
