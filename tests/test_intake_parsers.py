from __future__ import annotations

import importlib.util
import io
import json
import zipfile
from pathlib import Path

import pytest
from PIL import Image

from autoppt_service.parsers import IntakeLimits, SourceRejected, build_source_manifest
from autoppt_service.parsers.archive import extract_zip, inspect_zip
from autoppt_service.parsers.base import chunk_text, sha256_file
from autoppt_service.parsers.manifest import main as manifest_main
from autoppt_workflow.presentation.schemas import (
    EVIDENCE_ID_RE,
    normalize_evidence_item,
    normalize_evidence_ledger,
)


# ---------------------------------------------------------------- fixtures


def make_png(path: Path, size=(64, 32), colour="navy") -> Path:
    Image.new("RGB", size, colour).save(path, "PNG")
    return path


def make_text_pdf(path: Path, pages: list[str]) -> Path:
    """Minimal PDF with one Helvetica text line per page ('' = page without text)."""
    objects: dict[int, bytes] = {
        1: b"<< /Type /Catalog /Pages 2 0 R >>",
        3: b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    }
    kids = []
    for index, text in enumerate(pages):
        page_id, content_id = 4 + 2 * index, 5 + 2 * index
        kids.append(f"{page_id} 0 R")
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("latin-1") if text else b""
        objects[page_id] = (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {content_id} 0 R >>"
        ).encode()
        objects[content_id] = b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream"
    objects[2] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(pages)} >>".encode()
    out = bytearray(b"%PDF-1.4\n")
    offsets = {}
    for number in sorted(objects):
        offsets[number] = len(out)
        out += f"{number} 0 obj\n".encode() + objects[number] + b"\nendobj\n"
    xref = len(out)
    size = max(objects) + 1
    out += f"xref\n0 {size}\n0000000000 65535 f \n".encode()
    for number in range(1, size):
        out += f"{offsets[number]:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {size} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    path.write_bytes(bytes(out))
    return path


def make_scanned_pdf(path: Path) -> Path:
    Image.new("RGB", (200, 120), "white").save(path, "PDF")
    return path


def make_docx(path: Path, image: Path) -> Path:
    from docx import Document

    document = Document()
    document.add_heading("研究背景", level=1)
    document.add_paragraph("城市内涝频发，现有预警只能给出区县级结果。")
    document.add_heading("研究方法", level=1)
    document.add_paragraph("融合雷达降水与管网模型，输出街区级风险。")
    table = document.add_table(rows=3, cols=2)
    for row, values in zip(table.rows, [("方法", "准确率"), ("基线", "0.71"), ("本文", "0.86")]):
        for cell, value in zip(row.cells, values):
            cell.text = value
    document.add_picture(str(image))
    document.save(path)
    return path


def make_pptx(path: Path, image: Path) -> Path:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "实验结果"
    box = slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(4), Inches(1))
    box.text_frame.text = "本文方法在三个数据集上均优于基线"
    rows = slide.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(3), Inches(1)).table
    for r, values in enumerate([("数据集", "F1"), ("A", "0.9")]):
        for c, value in enumerate(values):
            rows.cell(r, c).text = value
    slide.shapes.add_picture(str(image), Inches(5), Inches(3))
    data = CategoryChartData()
    data.categories = ["A", "B", "C"]
    data.add_series("F1", (0.9, 0.8, 0.85))
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(5), Inches(1), Inches(4), Inches(2), data)
    slide.notes_slide.notes_text_frame.text = "强调数据集 C 的提升"
    presentation.slides.add_slide(presentation.slide_layouts[6])  # blank slide
    presentation.save(path)
    return path


def make_xlsx(path: Path) -> Path:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "结果"
    sheet.append(["方法", "准确率", "召回率"])
    sheet.append(["基线", 0.71, 0.65])
    sheet.append(["本文", 0.86, 0.80])
    notes = workbook.create_sheet("说明")
    notes.append(["备注"])
    notes.append(["数据来自 2026 年汛期"])
    workbook.save(path)
    return path


def entry(file_id: str, role: str, path: Path, **extra) -> dict:
    return {"file_id": file_id, "role": role, "name": path.name, "path": str(path), **extra}


def by_id(manifest: dict, file_id: str) -> dict:
    return next(record for record in manifest["files"] if record["file_id"] == file_id)


@pytest.fixture()
def sources(tmp_path: Path) -> dict[str, Path]:
    src = tmp_path / "src"
    src.mkdir()
    image = make_png(src / "图3-2.png")
    paths = {
        "image": image,
        "docx": make_docx(src / "论文终稿.docx", image),
        "pptx": make_pptx(src / "组会.pptx", image),
        "xlsx": make_xlsx(src / "实验结果.xlsx"),
        "pdf": make_text_pdf(src / "paper.pdf", ["Research question and method overview for flood alerts", ""]),
        "scan": make_scanned_pdf(src / "scan.pdf"),
    }
    (src / "data.csv").write_bytes("方法,准确率\n基线,0.71\n本文,0.86\n".encode("gb18030"))
    (src / "notes.md").write_text("# 背景\n内涝预警的现实问题。\n\n## 方法\n雷达与管网融合。\n", encoding="utf-8")
    (src / "abstract.txt").write_text("这是一段摘要文字。", encoding="utf-8")
    (src / "thesis.tex").write_text(
        "\\section{Introduction}\nFlood \\textbf{alerts} matter. % comment\n"
        "\\includegraphics[width=3in]{fig1.png}\n\\section{Method}\nRadar \\cite{x} fusion.\n",
        encoding="utf-8",
    )
    paths.update(csv=src / "data.csv", md=src / "notes.md", txt=src / "abstract.txt", tex=src / "thesis.tex")
    return paths


# ---------------------------------------------------------------- per format


def test_pdf_text_layer_and_ocr_flags(sources, tmp_path):
    manifest = build_source_manifest(
        [entry("f_01", "primary_source", sources["pdf"]), entry("f_02", "supplementary", sources["scan"])],
        tmp_path / "work",
    )
    text_pdf, scan_pdf = by_id(manifest, "f_01"), by_id(manifest, "f_02")
    assert text_pdf["stats"]["pages"] == 2
    assert text_pdf["status"] == "needs_ocr"  # page 2 has no text layer
    assert text_pdf["stats"]["ocr_pages"] == [2]
    page_items = [item for item in manifest["evidence_items"] if item["file_id"] == "f_01"]
    assert page_items[0]["page"] == 1
    assert "Research question" in page_items[0]["fact_or_data"]

    assert scan_pdf["status"] == "needs_ocr"
    queue = {(q["file_id"], q["page"]): q["reason"] for q in manifest["ocr_queue"]}
    assert queue[("f_01", 2)] == "no_text_layer"
    assert queue[("f_02", 1)] == "scanned_page"
    assert any(asset["file_id"] == "f_02" for asset in manifest["assets"])  # embedded scan image extracted


def test_pdf_page_limit(sources, tmp_path):
    manifest = build_source_manifest(
        [entry("f_01", "primary_source", sources["pdf"])], tmp_path / "work", limits=IntakeLimits(max_pdf_pages=1)
    )
    assert by_id(manifest, "f_01")["error"]["code"] == "pdf_too_many_pages"


def test_pptx_text_tables_chart_pictures_notes(sources, tmp_path):
    manifest = build_source_manifest([entry("f_01", "style_reference", sources["pptx"])], tmp_path / "work")
    record = by_id(manifest, "f_01")
    assert record["status"] == "parsed", record
    assert record["stats"]["slides"] == 2
    items = manifest["evidence_items"]
    slide_text = next(item for item in items if item["item_type"] == "text")
    assert slide_text["title"] == "实验结果"
    assert "优于基线" in slide_text["fact_or_data"]
    assert any(item["item_type"] == "speaker_notes" and "数据集 C" in item["fact_or_data"] for item in items)
    chart = next(table for table in manifest["tables"] if "from_chart" in table)
    assert chart["header"] == ["category", "F1"]
    assert chart["rows"][0] == ["A", 0.9]
    assert chart["chart_candidate"] is True
    assert any(table["header"] == ["数据集", "F1"] for table in manifest["tables"])
    assert manifest["assets"][0]["width"] == 64


def test_docx_reuses_repository_extractor(sources, tmp_path):
    manifest = build_source_manifest([entry("f_01", "primary_source", sources["docx"])], tmp_path / "work")
    record = by_id(manifest, "f_01")
    assert record["parser"] == "docx"
    assert record["stats"]["tables"] == 1
    chapters = [item["chapter"] for item in manifest["evidence_items"] if item["item_type"] == "text"]
    assert chapters == ["研究背景", "研究方法"]
    table = manifest["tables"][0]
    assert table["header"] == ["方法", "准确率"]
    assert table["numeric_columns"] == [1]
    assert manifest["assets"][0]["format"] == "PNG"


def test_xlsx_and_csv_tables(sources, tmp_path):
    manifest = build_source_manifest(
        [entry("f_01", "chart_data", sources["xlsx"]), entry("f_02", "chart_data", sources["csv"])],
        tmp_path / "work",
    )
    sheets = [table for table in manifest["tables"] if table["file_id"] == "f_01"]
    assert [table["sheet"] for table in sheets] == ["结果", "说明"]
    assert sheets[0]["numeric_columns"] == [1, 2]
    assert sheets[0]["chart_candidate"] is True
    assert sheets[1]["chart_candidate"] is False
    csv_table = next(table for table in manifest["tables"] if table["file_id"] == "f_02")
    assert by_id(manifest, "f_02")["stats"]["encoding"] == "gb18030"
    assert csv_table["header"] == ["方法", "准确率"]
    assert csv_table["rows"] == [["基线", "0.71"], ["本文", "0.86"]]
    assert manifest["summary"]["chart_candidates"] == [sheets[0]["table_id"], csv_table["table_id"]]


def test_sheet_row_limit_truncates(sources, tmp_path):
    manifest = build_source_manifest(
        [entry("f_01", "chart_data", sources["xlsx"])], tmp_path / "work", limits=IntakeLimits(max_sheet_rows=1)
    )
    assert manifest["tables"][0]["truncated"] is True
    assert manifest["tables"][0]["row_count"] == 1
    assert by_id(manifest, "f_01")["warnings"]


def test_markdown_txt_tex(sources, tmp_path):
    manifest = build_source_manifest(
        [
            entry("f_01", "supplementary", sources["md"]),
            entry("f_02", "supplementary", sources["txt"]),
            entry("f_03", "primary_source", sources["tex"]),
        ],
        tmp_path / "work",
    )
    items = manifest["evidence_items"]
    assert [i["chapter"] for i in items if i["file_id"] == "f_01"] == ["背景", "方法"]
    assert [i["fact_or_data"] for i in items if i["file_id"] == "f_02"] == ["这是一段摘要文字。"]
    tex_items = [i for i in items if i["file_id"] == "f_03"]
    assert [i["chapter"] for i in tex_items] == ["Introduction", "Method"]
    assert "Flood alerts matter." in tex_items[0]["fact_or_data"]
    assert "comment" not in tex_items[0]["fact_or_data"]
    assert "\\cite" not in tex_items[1]["fact_or_data"]
    assert by_id(manifest, "f_03")["stats"]["tex_graphics"] == ["fig1.png"]


def test_images_keep_as_image_and_rejections(sources, tmp_path):
    bad = tmp_path / "fake.png"
    bad.write_bytes(b"not an image")
    heic = tmp_path / "photo.heic"
    heic.write_bytes(b"\x00\x00\x00\x18ftypheic")
    big = make_png(tmp_path / "big.png", size=(200, 200))
    manifest = build_source_manifest(
        [
            entry("f_01", "keep_as_image", sources["image"]),
            entry("f_02", "image_asset", bad),
            entry("f_03", "image_asset", heic),
            entry("f_04", "image_asset", big),
        ],
        tmp_path / "work",
        limits=IntakeLimits(max_image_pixels=10_000),
    )
    asset = manifest["assets"][0]
    assert asset["keep_as_image"] is True
    assert manifest["summary"]["keep_as_image_assets"] == [asset["asset_id"]]
    figure = next(item for item in manifest["evidence_items"] if item["file_id"] == "f_01")
    assert figure["item_type"] == "figure"
    assert figure["figure"] == asset["asset_id"]
    assert by_id(manifest, "f_02")["error"]["code"] == "image_unreadable"
    if importlib.util.find_spec("pillow_heif") is None:
        assert by_id(manifest, "f_03")["status"] == "needs_transcode"
    assert by_id(manifest, "f_04")["error"]["code"] == "image_too_large"


# ---------------------------------------------------------------- ZIP safety


def write_zip(path: Path, members: dict[str, bytes], *, compression=zipfile.ZIP_DEFLATED) -> Path:
    with zipfile.ZipFile(path, "w", compression) as archive:
        for name, data in members.items():
            archive.writestr(name, data)
    return path


def png_bytes() -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (8, 8), "red").save(buffer, "PNG")
    return buffer.getvalue()


def test_zip_dispatches_members_by_extension(tmp_path):
    archive = write_zip(
        tmp_path / "资料.zip",
        {
            "docs/摘要.md": "# 摘要\n内容".encode("utf-8"),
            "docs/图1.png": png_bytes(),
            "__MACOSX/docs/._图1.png": b"junk",
            "tool.exe": b"MZ",
            "inner.zip": b"PK",
        },
    )
    manifest = build_source_manifest([entry("f_05", "primary_source", archive)], tmp_path / "work")
    parent = by_id(manifest, "f_05")
    assert parent["status"] == "unpacked"
    reasons = {item["member"]: item["reason"] for item in parent["stats"]["skipped"]}
    assert reasons == {
        "__MACOSX/docs/._图1.png": "system_file",
        "tool.exe": "unsupported_type",
        "inner.zip": "nested_archive",
        "docs/图1.png": "role_mismatch",  # images are not primary sources
    }
    child = by_id(manifest, parent["stats"]["children"][0])
    assert child["parent_file_id"] == "f_05"
    assert child["archive_member"] == "docs/摘要.md"
    assert child["status"] == "parsed"
    item = next(i for i in manifest["evidence_items"] if i["file_id"] == child["file_id"])
    assert item["archive_member"] == "docs/摘要.md"
    assert Path(child["path"]).resolve().is_relative_to((tmp_path / "work" / "unpacked" / "f_05").resolve())


@pytest.mark.parametrize("name", ["../evil.md", "/abs/evil.md", "C:/evil.md", "a/../../evil.md", "..\\evil.md"])
def test_zip_rejects_path_traversal(tmp_path, name):
    archive = write_zip(tmp_path / "bad.zip", {"ok.md": b"# ok", name: b"x"})
    manifest = build_source_manifest([entry("f_01", "supplementary", archive)], tmp_path / "work")
    assert by_id(manifest, "f_01")["error"]["code"] == "zip_path_traversal"
    assert not (tmp_path / "evil.md").exists()
    assert not (tmp_path / "work" / "unpacked").exists()


def test_zip_rejects_bomb_by_ratio(tmp_path):
    archive = write_zip(tmp_path / "bomb.zip", {"zeros.txt": b"\0" * (2 * 1024 * 1024)})
    with pytest.raises(SourceRejected) as caught:
        inspect_zip(archive, IntakeLimits())
    assert caught.value.code == "zip_bomb"


def test_zip_size_and_count_limits(tmp_path):
    archive = write_zip(tmp_path / "many.zip", {f"{i}.md": b"# x" for i in range(5)}, compression=zipfile.ZIP_STORED)
    with pytest.raises(SourceRejected) as caught:
        inspect_zip(archive, IntakeLimits(max_zip_entries=4))
    assert caught.value.code == "zip_too_many_entries"
    with pytest.raises(SourceRejected) as caught:
        inspect_zip(archive, IntakeLimits(max_zip_total_bytes=10))
    assert caught.value.code == "zip_total_too_large"
    with pytest.raises(SourceRejected) as caught:
        inspect_zip(archive, IntakeLimits(max_zip_entry_bytes=2))
    assert caught.value.code == "zip_entry_too_large"


def test_zip_rejects_symlink_and_lies_about_size(tmp_path):
    archive = tmp_path / "link.zip"
    with zipfile.ZipFile(archive, "w") as handle:
        info = zipfile.ZipInfo("link.md")
        info.external_attr = (0o120777 << 16)
        handle.writestr(info, "/etc/passwd")
    with pytest.raises(SourceRejected) as caught:
        inspect_zip(archive, IntakeLimits())
    assert caught.value.code == "zip_symlink"

    # A header that under-declares the size must not let extraction write more.
    liar = write_zip(tmp_path / "liar.zip", {"a.md": b"# " + b"x" * 5000}, compression=zipfile.ZIP_STORED)
    data = bytearray(liar.read_bytes())
    real = (5002).to_bytes(4, "little")
    fake = (10).to_bytes(4, "little")
    data = data.replace(real + real, real + fake)  # local header: compressed, uncompressed
    data = data.replace(real + real, real + fake)  # central directory
    liar.write_bytes(bytes(data))
    with pytest.raises(SourceRejected) as caught:
        extract_zip(liar, tmp_path / "out", IntakeLimits())
    assert caught.value.code in {"bad_zip", "zip_bomb"}
    assert not (tmp_path / "out").exists()


def test_gbk_member_names_are_decoded(tmp_path):
    # Chinese Windows archivers store GBK bytes without the UTF-8 flag.
    gbk = "摘要.md".encode("gbk")
    placeholder = b"X" * (len(gbk) - 3) + b".md"
    archive = write_zip(tmp_path / "gbk.zip", {placeholder.decode("ascii"): "# 摘要".encode("utf-8")})
    archive.write_bytes(archive.read_bytes().replace(placeholder, gbk))
    unpacked = extract_zip(archive, tmp_path / "out", IntakeLimits())
    assert unpacked["files"][0]["member"] == "摘要.md"
    assert (tmp_path / "out" / "摘要.md").is_file()


def test_ooxml_bomb_is_rejected_before_parsing(tmp_path):
    fake_docx = write_zip(tmp_path / "bomb.docx", {"word/document.xml": b" " * (2 * 1024 * 1024)})
    manifest = build_source_manifest([entry("f_01", "primary_source", fake_docx)], tmp_path / "work")
    assert by_id(manifest, "f_01")["error"]["code"] == "zip_bomb"


# ---------------------------------------------------------------- manifest


def test_manifest_contract_and_evidence_normalization(sources, tmp_path):
    files = [
        entry("f_01", "primary_source", sources["docx"], sha256=sha256_file(sources["docx"])),
        entry("f_02", "chart_data", sources["xlsx"]),
        entry("f_03", "keep_as_image", sources["image"]),
        entry("f_04", "style_reference", sources["pptx"]),
        entry("f_05", "supplementary", sources["pdf"]),
        entry("f_06", "supplementary", sources["md"]),
    ]
    manifest = build_source_manifest(files, tmp_path / "work", description="重点讲清本人贡献。")
    assert manifest["schema_version"] == "autoppt-source-manifest-v1"
    assert manifest["errors"] == []
    assert manifest["summary"]["rejected"] == []
    json.dumps(manifest, ensure_ascii=False)

    user = manifest["evidence_items"][0]
    assert user["item_type"] == "user_statement"
    assert user["reliability"] == "user_statement"

    ids = [item["evidence_id"] for item in manifest["evidence_items"]]
    assert len(ids) == len(set(ids))
    for index, item in enumerate(manifest["evidence_items"], 1):
        normalized = normalize_evidence_item(item, index=index)
        assert normalized["evidence_id"] == item["evidence_id"]
        assert EVIDENCE_ID_RE.fullmatch(normalized["evidence_id"])
        assert normalized["source_file"]
        assert normalized["fact_or_data"]
        assert normalized["confirmed"] is False
    ledger = normalize_evidence_ledger([manifest], source_label="intake")
    assert ledger["summary"]["count"] == len(ids)
    assert ledger["summary"]["confirmed_count"] == 0


def test_manifest_rejects_hash_mismatch_missing_and_role_mismatch(sources, tmp_path):
    tool = tmp_path / "tool.exe"
    tool.write_bytes(b"MZ")
    manifest = build_source_manifest(
        [
            entry("f_01", "primary_source", sources["docx"], sha256="0" * 64),
            entry("f_02", "primary_source", tmp_path / "missing.pdf"),
            entry("f_03", "chart_data", sources["md"]),
            entry("f_04", "supplementary", tool),
        ],
        tmp_path / "work",
    )
    assert by_id(manifest, "f_01")["error"]["code"] == "sha256_mismatch"
    assert by_id(manifest, "f_02")["error"]["code"] == "file_missing"
    assert by_id(manifest, "f_03")["error"]["code"] == "file_role_mismatch"
    assert by_id(manifest, "f_04")["error"]["code"] == "unsupported_file_type"
    assert len(manifest["summary"]["rejected"]) == 4


def test_manifest_request_level_limits(sources, tmp_path):
    files = [entry("f_01", "supplementary", sources["md"]), entry("f_02", "supplementary", sources["txt"])]
    manifest = build_source_manifest(files, tmp_path / "work", limits=IntakeLimits(max_files=1, max_total_bytes=1))
    assert {error["code"] for error in manifest["errors"]} == {"too_many_files", "total_too_large"}
    assert manifest["files"] == []


def test_long_text_is_chunked(tmp_path):
    long_text = tmp_path / "long.txt"
    long_text.write_text(("句子内容很长。" * 100 + "\n\n") * 5, encoding="utf-8")
    manifest = build_source_manifest(
        [entry("f_01", "supplementary", long_text)], tmp_path / "work", limits=IntakeLimits(max_item_chars=500)
    )
    items = manifest["evidence_items"]
    assert len(items) > 1
    assert all(len(item["fact_or_data"]) <= 500 for item in items)
    assert chunk_text("a" * 1200, 500) == ["a" * 500, "a" * 500, "a" * 200]


def test_manifest_cli(sources, tmp_path):
    listing = tmp_path / "files.json"
    listing.write_text(json.dumps([entry("f_01", "supplementary", sources["md"])], ensure_ascii=False), encoding="utf-8")
    output = tmp_path / "manifest.json"
    assert manifest_main([str(listing), "--work-dir", str(tmp_path / "work"), "--output", str(output)]) == 0
    assert json.loads(output.read_text(encoding="utf-8"))["summary"]["files"] == 1
