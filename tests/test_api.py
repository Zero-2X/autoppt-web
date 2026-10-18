"""Backend API with a mock Worker: full job life cycle, credits, files and auth."""

from __future__ import annotations

import hashlib
import io
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from autoppt_api.app import create_app
from autoppt_api.config import Settings
from autoppt_api import ledger
from autoppt_service.worker_client import ApiWorker, restore_snapshot, snapshot_bytes

DESCRIPTION = "本研究面向城市内涝预警，目标是在暴雨前两小时给出街区级风险等级，并说明本人负责的模型部分。" * 6


@pytest.fixture()
def api(tmp_path):
    settings = Settings(
        database_url=f"sqlite:///{(tmp_path / 'api.db').as_posix()}",
        storage_dir=tmp_path / "storage",
        dev_echo_codes=True,
        admin_emails=("admin@example.com",),
    )
    client = TestClient(create_app(settings))
    client.tmp = tmp_path
    return client


def login(client: TestClient, email: str) -> dict[str, str]:
    code = client.post("/api/auth/code", json={"email": email}).json()["dev_code"]
    tokens = client.post("/api/auth/login", json={"email": email, "code": code}).json()
    return {"Authorization": f"Bearer {tokens['access_token']}"}


def upload(client: TestClient, headers: dict, name: str, data: bytes, chunk: int = 4096) -> str:
    init = client.post("/api/uploads", json={"filename": name, "size": len(data)}, headers=headers).json()
    for offset in range(0, len(data), chunk):
        reply = client.put(init["upload_url"], params={"offset": offset}, content=data[offset:offset + chunk])
        assert reply.status_code == 200, reply.text
    done = client.post(f"/api/uploads/{init['upload_id']}/complete", headers=headers)
    assert done.status_code == 200, done.text
    return done.json()["sha256"]


def docx_bytes() -> bytes:
    from docx import Document

    document = Document()
    document.add_heading("研究背景", level=1)
    document.add_paragraph("城市内涝频发，现有预警只能给出区县级结果，平均提前 30 分钟。")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def request_body(sha: str, tier: str = "native_editable") -> dict:
    delivery = {"tier": tier, "font_policy": "office_native_yahei"} if tier == "native_editable" else {"tier": tier}
    return {
        "schema_version": "autoppt-job-request-v1",
        "service": "new_deck",
        "inputs": {"files": [{"file_id": "f_01", "role": "primary_source", "name": "论文.docx", "sha256": sha}],
                   "description": DESCRIPTION},
        "scenario": "thesis_defense",
        "slides": {"mode": "fixed", "count": 3},
        "cover": {"title": "城市内涝预警研究"},
        "delivery": delivery,
    }


@pytest.fixture()
def world(api):
    admin = login(api, "admin@example.com")
    user = login(api, "student@example.com")
    worker = api.post("/api/admin/workers", json={"name": "w1", "capabilities": ["all"]}, headers=admin).json()
    api.post("/api/admin/credits/adjust", json={"email": "student@example.com", "credits": 500, "memo": "beta"},
             headers={**admin, "Idempotency-Key": "beta-1"})
    sha = upload(api, user, "论文.docx", docx_bytes())
    mock_worker = ApiWorker(api, worker["token"], api.tmp / "worker-jobs", mock=True, heartbeat_s=3600, min_free_gb=0)
    return {"admin": admin, "user": user, "worker": mock_worker, "worker_token": worker["token"], "sha": sha}


def submit(api, world, tier="native_editable") -> dict:
    job = api.post("/api/jobs", json={"request": request_body(world["sha"], tier)}, headers=world["user"]).json()
    estimate = api.post(f"/api/jobs/{job['id']}/estimate", headers=world["user"])
    assert estimate.status_code == 200, estimate.text
    total = estimate.json()["total"]
    reply = api.post(f"/api/jobs/{job['id']}/submit", json={"confirm_total": total}, headers={**world["user"], "Idempotency-Key": "s1"})
    assert reply.status_code == 200, reply.text
    return reply.json()


def credits(api, headers) -> dict:
    return api.get("/api/credits", headers=headers).json()


def ledger_ok(api) -> None:
    with api.app.state.session_factory() as session:
        assert ledger.check_ledger(session) == []


# ---------------------------------------------------------------- full life cycle


def test_full_native_job_with_three_checkpoints(api, world):
    user, worker = world["user"], world["worker"]
    job = submit(api, world)
    assert job["state"] == "queued" and job["hold"]["amount"] == job["estimate"]["total"]
    before = credits(api, user)
    assert before["frozen"] == job["estimate"]["total"]

    assert worker.run_once()["job_state"] == "planning"       # intake
    assert worker.run_once()["job_state"] == "cp1_outline"    # planning
    assert worker.run_once() is None                          # waiting for the user
    cp1 = api.get(f"/api/jobs/{job['id']}/checkpoint", headers=user).json()
    assert cp1["type"] == "outline" and cp1["previews"][0]["url"]
    assert api.get(cp1["previews"][0]["url"]).status_code == 200
    reply = api.post(f"/api/jobs/{job['id']}/checkpoint/{cp1['checkpoint_id']}/response",
                     json={"decision": "approve", "user_reply": "大纲可以"}, headers=user)
    assert reply.status_code == 200, reply.text

    assert worker.run_once()["job_state"] == "cp2_masters"
    cp2 = api.get(f"/api/jobs/{job['id']}/checkpoint", headers=user).json()
    redo = {"decision": "redo_slides", "user_reply": "第一页重做", "redo_slide_ids": ["S01"]}
    unconfirmed = api.post(f"/api/jobs/{job['id']}/checkpoint/{cp2['checkpoint_id']}/response", json=redo, headers=user)
    assert unconfirmed.status_code == 422 and unconfirmed.json()["details"]["charge"] == cp2["pricing"]["redo_slide"]
    confirmed = api.post(f"/api/jobs/{job['id']}/checkpoint/{cp2['checkpoint_id']}/response",
                         json={**redo, "charge_confirmed": True}, headers=user)
    assert confirmed.status_code == 200, confirmed.text
    assert worker.run_once()["job_state"] == "cp2_masters"   # redone, round 2
    cp2b = api.get(f"/api/jobs/{job['id']}/checkpoint", headers=user).json()
    assert cp2b["round"] == 2
    api.post(f"/api/jobs/{job['id']}/checkpoint/{cp2b['checkpoint_id']}/response",
             json={"decision": "approve", "user_reply": "通过"}, headers=user)

    assert worker.run_once()["job_state"] == "cp3_native_preview"
    cp3 = api.get(f"/api/jobs/{job['id']}/checkpoint", headers=user).json()
    assert cp3["editability_scope"]
    side = [p["path"] for p in cp3["previews"] if p["kind"] == "side_by_side"]
    # An administrator can never approve CP3 on the user's behalf.
    admin_body = {"decision": "approve", "user_reply": "ok", "reviewed_previews": side}
    url = f"/api/jobs/{job['id']}/checkpoint/{cp3['checkpoint_id']}/response"
    assert api.post(url, json=admin_body, headers=world["admin"]).status_code == 403  # no access reason
    admin_try = api.post(url, json=admin_body, headers={**world["admin"], "X-Access-Reason": "support ticket 42"})
    assert admin_try.status_code == 422  # even with a reason, an admin never approves CP3
    audit = api.get("/api/admin/audit", headers=world["admin"]).json()
    assert any(a["action"] == "admin.access_job" and a["data"]["reason"] == "support ticket 42" for a in audit)
    api.post(f"/api/jobs/{job['id']}/checkpoint/{cp3['checkpoint_id']}/response",
             json={"decision": "approve", "user_reply": "接受预览中的非文字差异，优先真实可编辑", "reviewed_previews": side},
             headers=user)

    assert worker.run_once()["job_state"] == "delivered"
    detail = api.get(f"/api/jobs/{job['id']}", headers=user).json()
    assert detail["state"] == "delivered" and detail["hold"]["status"] == "closed"
    paths = {d["path"] for d in detail["deliverables"]}
    assert "output/native-deck.pptx" in paths and "service/acceptance.json" in paths
    assert api.get(detail["deliverables"][0]["url"]).status_code == 200
    after = credits(api, user)
    phases = job["estimate"]["phases"]
    charged = phases["planning"] + phases["masters"] + phases["reconstruct"] + phases["package"] + cp2["pricing"]["redo_slide"]
    assert before["balance"] - after["balance"] == charged
    assert after["frozen"] == 0
    kinds = [e["kind"] for e in api.get(f"/api/jobs/{job['id']}/events", headers=user).json()]
    assert kinds.count("checkpoint") == 4 and "delivered" in kinds
    ledger_ok(api)


def test_cancel_releases_frozen_credits(api, world):
    job = submit(api, world, "image_deck")
    world["worker"].run_once()
    world["worker"].run_once()  # planning settled, waiting at CP1
    before = credits(api, world["user"])
    reply = api.post(f"/api/jobs/{job['id']}/cancel", headers=world["user"]).json()
    assert reply["state"] == "cancelled"
    after = credits(api, world["user"])
    assert after["frozen"] == 0 and after["balance"] == before["balance"]
    assert after["available"] == before["available"] + before["frozen"]
    assert world["worker"].run_once() is None
    ledger_ok(api)


def test_system_failure_refunds_unfinished_work(api, world):
    job = submit(api, world)
    lease = api.post("/api/worker/lease", json={}, headers={"Authorization": f"Bearer {world['worker_token']}"}).json()
    failed = api.post(f"/api/worker/jobs/{job['id']}/fail", json={"lease_token": lease["lease_token"], "error": "disk full",
                                                                  "retryable": False},
                      headers={"Authorization": f"Bearer {world['worker_token']}"}).json()
    assert failed["state"] == "failed"
    assert credits(api, world["user"])["frozen"] == 0
    ledger_ok(api)


def test_retryable_failure_backs_off_then_reruns(api, world):
    job = submit(api, world)
    token = {"Authorization": f"Bearer {world['worker_token']}"}
    lease = api.post("/api/worker/lease", json={}, headers=token).json()
    reply = api.post(f"/api/worker/jobs/{job['id']}/fail", json={"lease_token": lease["lease_token"], "error": "COM busy"}, headers=token)
    assert reply.json()["state"] == "queued"
    assert api.post("/api/worker/lease", json={}, headers=token).status_code == 204  # backing off


def test_leases_are_exclusive_and_expire(api, world):
    admin = world["admin"]
    other = api.post("/api/admin/workers", json={"name": "w2", "capabilities": ["all"]}, headers=admin).json()
    job = submit(api, world)
    first = api.post("/api/worker/lease", json={}, headers={"Authorization": f"Bearer {world['worker_token']}"}).json()
    assert first["job_id"] == job["id"]
    assert api.post("/api/worker/lease", json={}, headers={"Authorization": f"Bearer {other['token']}"}).status_code == 204
    # Expire the lease: the second worker takes over and the first worker's report is refused.
    from autoppt_api.models import Job
    from datetime import datetime, timedelta, timezone

    with api.app.state.session_factory() as session:
        row = session.get(Job, job["id"])
        row.lease_expires_at = datetime.now(timezone.utc) - timedelta(seconds=1)
        session.commit()
    second = api.post("/api/worker/lease", json={}, headers={"Authorization": f"Bearer {other['token']}"}).json()
    assert second["job_id"] == job["id"] and second["lease_token"] != first["lease_token"]
    stale = api.post(f"/api/worker/jobs/{job['id']}/complete",
                     json={"lease_token": first["lease_token"], "phase": "intake", "job_state": "planning"},
                     headers={"Authorization": f"Bearer {world['worker_token']}"})
    assert stale.status_code == 409


def test_capabilities_route_phases(api, world):
    admin = world["admin"]
    office_only = api.post("/api/admin/workers", json={"name": "office", "capabilities": ["office"]}, headers=admin).json()
    submit(api, world)
    assert api.post("/api/worker/lease", json={}, headers={"Authorization": f"Bearer {office_only['token']}"}).status_code == 204


# ---------------------------------------------------------------- orders and credits


def test_insufficient_credits_and_price_confirmation(api, world):
    poor = login(api, "poor@example.com")
    sha = upload(api, poor, "论文.docx", docx_bytes())
    job = api.post("/api/jobs", json={"request": request_body(sha)}, headers=poor).json()
    total = api.post(f"/api/jobs/{job['id']}/estimate", headers=poor).json()["total"]
    assert api.post(f"/api/jobs/{job['id']}/submit", json={"confirm_total": total + 1}, headers=poor).status_code == 409
    reply = api.post(f"/api/jobs/{job['id']}/submit", json={"confirm_total": total}, headers=poor)
    assert reply.status_code == 402 and reply.json()["details"]["available"] == 0


def test_submit_is_idempotent(api, world):
    job = api.post("/api/jobs", json={"request": request_body(world["sha"])}, headers=world["user"]).json()
    total = api.post(f"/api/jobs/{job['id']}/estimate", headers=world["user"]).json()["total"]
    headers = {**world["user"], "Idempotency-Key": "same"}
    first = api.post(f"/api/jobs/{job['id']}/submit", json={"confirm_total": total}, headers=headers).json()
    second = api.post(f"/api/jobs/{job['id']}/submit", json={"confirm_total": total}, headers=headers).json()
    assert first["state"] == second["state"] == "queued"
    assert credits(api, world["user"])["frozen"] == total
    ledger_ok(api)


def test_redeem_codes_and_admin_adjust(api, world):
    codes = api.post("/api/admin/redeem-codes", json={"credits": 50, "count": 2}, headers=world["admin"]).json()["codes"]
    before = credits(api, world["user"])["balance"]
    assert api.post("/api/credits/redeem", json={"code": codes[0]}, headers=world["user"]).json()["credited"] == 50
    assert api.post("/api/credits/redeem", json={"code": codes[0]}, headers=world["user"]).json()["credited"] == 50  # same posting
    assert credits(api, world["user"])["balance"] == before + 50
    other = login(api, "other@example.com")
    assert api.post("/api/credits/redeem", json={"code": codes[0]}, headers=other).status_code == 409  # single use
    assert api.post("/api/credits/redeem", json={"code": "NOPE"}, headers=other).status_code == 404
    # Repeating an adjustment with the same Idempotency-Key does not credit twice.
    api.post("/api/admin/credits/adjust", json={"email": "student@example.com", "credits": 500, "memo": "beta"},
             headers={**world["admin"], "Idempotency-Key": "beta-1"})
    assert credits(api, world["user"])["balance"] == before + 50
    assert api.post("/api/admin/redeem-codes", json={"credits": 5}, headers=world["user"]).status_code == 403
    assert api.get("/api/admin/ledger/check", headers=world["admin"]).json()["ok"] is True


# ---------------------------------------------------------------- files


def test_uploads_sniff_content_and_deduplicate(api, world):
    user = world["user"]
    fake = upload(api, user, "fake.docx", "这其实是文本".encode("utf-8"))
    job = api.post("/api/jobs", json={"request": request_body(fake)}, headers=user).json()
    reply = api.post(f"/api/jobs/{job['id']}/estimate", headers=user)
    assert reply.status_code == 422
    assert "file_type_mismatch" in {e["code"] for e in reply.json()["details"]["errors"]}
    # A known hash never completes an upload without the bytes (proof of possession).
    data = docx_bytes()
    init = api.post("/api/uploads", json={"filename": "again.docx", "size": len(data), "sha256": world["sha"]}, headers=user).json()
    assert init["complete"] is False and init["upload_url"]
    # Someone else's hash cannot be referenced without uploading it.
    other = login(api, "other@example.com")
    stolen = api.post("/api/jobs", json={"request": request_body(world["sha"])}, headers=other).json()
    errors = api.post(f"/api/jobs/{stolen['id']}/estimate", headers=other).json()["details"]["errors"]
    assert "upload_not_found" in {e["code"] for e in errors}


def test_upload_rejects_bad_offsets_tokens_and_size(api, world):
    user = world["user"]
    init = api.post("/api/uploads", json={"filename": "a.txt", "size": 10}, headers=user).json()
    assert api.put(init["upload_url"], params={"offset": 5}, content=b"12345").status_code == 409
    assert api.put(init["upload_url"] + "x", params={"offset": 0}, content=b"1").status_code == 403
    assert api.put(init["upload_url"], params={"offset": 0}, content=b"0123456789ABC").status_code == 409
    assert api.post("/api/uploads", json={"filename": "big.pdf", "size": 10**12}, headers=user).status_code == 413


def test_delete_job_removes_files(api, world):
    user = world["user"]
    sha = upload(api, user, "单独.docx", docx_bytes() + b"")
    job = api.post("/api/jobs", json={"request": request_body(sha)}, headers=user).json()
    reply = api.delete(f"/api/jobs/{job['id']}", headers=user).json()
    assert reply["deleted"] is True
    assert api.get(f"/api/jobs/{job['id']}", headers=user).status_code == 404
    audit = api.get("/api/admin/audit", headers=world["admin"]).json()
    assert any(a["action"] == "job.delete" for a in audit)


# ---------------------------------------------------------------- auth


def test_auth_refresh_rotation_and_reuse_detection(api):
    code = api.post("/api/auth/code", json={"email": "a@example.com"}).json()["dev_code"]
    assert api.post("/api/auth/login", json={"email": "a@example.com", "code": "000000" if code != "000000" else "111111"}).status_code == 401
    code = api.post("/api/auth/code", json={"email": "a@example.com"}).json()["dev_code"]
    tokens = api.post("/api/auth/login", json={"email": "a@example.com", "code": code}).json()
    rotated = api.post("/api/auth/refresh", json={"refresh_token": tokens["refresh_token"]})
    assert rotated.status_code == 200
    # Reusing the old refresh token revokes the whole family, including the new one.
    assert api.post("/api/auth/refresh", json={"refresh_token": tokens["refresh_token"]}).status_code == 401
    assert api.post("/api/auth/refresh", json={"refresh_token": rotated.json()["refresh_token"]}).status_code == 401
    assert api.get("/api/me", headers={"Authorization": "Bearer nope"}).status_code == 401


def test_login_code_rate_limit_and_attempts(api):
    for _ in range(5):
        assert api.post("/api/auth/code", json={"email": "r@example.com"}).status_code == 200
    assert api.post("/api/auth/code", json={"email": "r@example.com"}).status_code == 429
    assert api.post("/api/auth/code", json={"email": "not-an-email"}).status_code == 422


def test_options_and_openapi(api):
    assert len(api.get("/api/options").json()["scenarios"]) == 6
    spec = api.get("/openapi.json").json()
    assert "/api/jobs/{job_id}/checkpoint/{checkpoint_id}/response" in spec["paths"]
    assert "/api/worker/lease" in spec["paths"]


def test_event_stream(api, world):
    job = submit(api, world)
    response = api.get(f"/api/jobs/{job['id']}/events/stream", params={"timeout": 0}, headers=world["user"])
    assert response.status_code == 200 and "event: state" in response.text and "event: end" in response.text


def test_snapshot_round_trip_rejects_escape(tmp_path):
    job = tmp_path / "job"
    (job / "service").mkdir(parents=True)
    (job / "service" / "a.json").write_text("{}", encoding="utf-8")
    data = snapshot_bytes(job)
    restore_snapshot(data, tmp_path / "copy")
    assert (tmp_path / "copy" / "service" / "a.json").read_text(encoding="utf-8") == "{}"
    import zipfile

    evil = io.BytesIO()
    with zipfile.ZipFile(evil, "w") as archive:
        archive.writestr("../escape.txt", "x")
    from autoppt_service.parsers.base import SourceRejected

    with pytest.raises(SourceRejected):
        restore_snapshot(evil.getvalue(), tmp_path / "copy2")
    assert not (tmp_path / "escape.txt").exists()


def test_retention_purge_and_notifications(api, world):
    from datetime import datetime, timedelta, timezone

    from autoppt_api.lifecycle import purge_expired
    from autoppt_api.maintenance import send_notifications
    from autoppt_api.models import Job, Notification, StoredFile

    job = submit(api, world, "planning_package")
    for _ in range(3):
        world["worker"].run_once()
    cp1 = api.get(f"/api/jobs/{job['id']}/checkpoint", headers=world["user"]).json()
    api.post(f"/api/jobs/{job['id']}/checkpoint/{cp1['checkpoint_id']}/response",
             json={"decision": "approve", "user_reply": "好"}, headers=world["user"])
    assert world["worker"].run_once()["job_state"] == "delivered"
    with api.app.state.session_factory() as session:
        assert send_notifications(session, "") >= 2  # checkpoint + delivered, logged without SMTP
        assert session.query(Notification).filter(Notification.status == "pending").count() == 0
        assert purge_expired(session, api.app.state.storage) == []  # still inside the retention period
        row = session.get(Job, job["id"])
        row.finished_at = datetime.now(timezone.utc) - timedelta(days=8)  # default retention: 7 days
        session.commit()
        assert purge_expired(session, api.app.state.storage) == [job["id"]]
        session.commit()
        assert session.get(Job, job["id"]).state == "deleted"
        assert session.get(StoredFile, world["sha"]) is None  # retention covers the uploaded source too
    assert api.get(f"/api/jobs/{job['id']}", headers=world["user"]).status_code == 404
    ledger_ok(api)


def test_alembic_migration_matches_models(tmp_path, monkeypatch):
    from alembic import command
    from alembic.config import Config

    monkeypatch.setenv("AUTOPPT_API_DATABASE_URL", f"sqlite:///{(tmp_path / 'm.db').as_posix()}")
    root = Path(__file__).resolve().parents[1]
    config = Config(str(root / "alembic.ini"))
    config.set_main_option("script_location", str(root / "autoppt_api" / "migrations"))
    command.upgrade(config, "head")
    command.check(config)  # raises if the models drifted from the migrations


def test_per_slide_phases_settle_actual_slides_not_the_frozen_maximum(api, world):
    user, worker = world["user"], world["worker"]
    body = request_body(world["sha"], "image_deck")
    body["slides"] = {"mode": "range", "min": 8, "max": 16}  # mock planning picks the middle: 12 slides
    job = api.post("/api/jobs", json={"request": body}, headers=user).json()
    quote = api.post(f"/api/jobs/{job['id']}/estimate", headers=user).json()
    assert quote["slides_for_hold"] == 16
    before = credits(api, user)
    api.post(f"/api/jobs/{job['id']}/submit", json={"confirm_total": quote["total"]}, headers=user)
    worker.run_once(), worker.run_once()
    for _ in range(2):
        cp = api.get(f"/api/jobs/{job['id']}/checkpoint", headers=user).json()
        api.post(f"/api/jobs/{job['id']}/checkpoint/{cp['checkpoint_id']}/response",
                 json={"decision": "approve", "user_reply": "通过"}, headers=user)
        assert worker.run_once()["job_state"] in ("cp2_masters", "delivered")
    assert api.get(f"/api/jobs/{job['id']}", headers=user).json()["state"] == "delivered"
    after = credits(api, user)
    expected = quote["phases"]["planning"] + quote["phases"]["masters"] * 12 // 16
    assert before["balance"] - after["balance"] == expected < quote["total"]
    assert after["frozen"] == 0
    ledger_ok(api)


def test_vectorize_job_through_the_api(api, world):
    from pptx import Presentation

    user, worker = world["user"], world["worker"]
    presentation = Presentation()
    for _ in range(2):
        presentation.slides.add_slide(presentation.slide_layouts[6])
    buffer = io.BytesIO()
    presentation.save(buffer)
    sha = upload(api, user, "旧稿.pptx", buffer.getvalue())
    body = {"schema_version": "autoppt-job-request-v1", "service": "vectorize",
            "inputs": {"files": [{"file_id": "f_01", "role": "vectorize_source", "name": "旧稿.pptx", "sha256": sha}]},
            "slides": {"mode": "fixed", "count": 2},
            "delivery": {"tier": "native_editable", "font_policy": "office_native_yahei"}}
    job = api.post("/api/jobs", json={"request": body}, headers=user).json()
    quote = api.post(f"/api/jobs/{job['id']}/estimate", headers=user).json()
    assert quote["phases"]["reconstruct"] == 2 * 12
    api.post(f"/api/jobs/{job['id']}/submit", json={"confirm_total": quote["total"]}, headers=user)
    assert worker.run_once()["job_state"] == "cp_v_text"
    cpv = api.get(f"/api/jobs/{job['id']}/checkpoint", headers=user).json()
    assert cpv["type"] == "vectorize_text" and any(p["kind"] == "document" for p in cpv["previews"])
    api.post(f"/api/jobs/{job['id']}/checkpoint/{cpv['checkpoint_id']}/response",
             json={"decision": "approve", "user_reply": "文字已核对", "answers": {"text_corrections": {"S01": {"line-000": "新标题"}}}},
             headers=user)
    assert worker.run_once()["job_state"] == "cp3_native_preview"
    cp3 = api.get(f"/api/jobs/{job['id']}/checkpoint", headers=user).json()
    side = [p["path"] for p in cp3["previews"] if p["kind"] == "side_by_side"]
    api.post(f"/api/jobs/{job['id']}/checkpoint/{cp3['checkpoint_id']}/response",
             json={"decision": "approve", "user_reply": "接受", "reviewed_previews": side}, headers=user)
    assert worker.run_once()["job_state"] == "delivered"
    detail = api.get(f"/api/jobs/{job['id']}", headers=user).json()
    acceptance_url = next(d["url"] for d in detail["deliverables"] if d["path"] == "service/acceptance.json")
    acceptance = api.get(acceptance_url).json()
    assert acceptance["master_provenance"] == "user_supplied_master"
    paths = {d["path"] for d in detail["deliverables"]}
    assert "workspace/vectorize/reviewed-text.json" in paths
    ledger_ok(api)


def test_admin_refunds_settled_credits_of_one_job(api, world):
    admin, user = world["admin"], world["user"]
    job = submit(api, world, "image_deck")
    world["worker"].run_once()
    world["worker"].run_once()  # planning settled, waiting at CP1
    settled = job["estimate"]["phases"]["planning"]
    before = credits(api, user)
    url = f"/api/admin/jobs/{job['id']}/refund"
    assert api.post(url, json={"credits": 1, "reason": "x"}, headers=user).status_code == 403
    too_much = api.post(url, json={"credits": settled + 1, "reason": "大纲质量差"}, headers=admin)
    assert too_much.status_code == 409 and str(settled) in too_much.text
    assert api.post(url, json={"credits": 1, "reason": ""}, headers=admin).status_code == 422

    once = {**admin, "Idempotency-Key": "r-1"}
    reply = api.post(url, json={"credits": settled, "reason": "大纲质量差"}, headers=once)
    assert reply.status_code == 200, reply.text
    assert reply.json()["refundable"] == 0
    assert api.post(url, json={"credits": settled, "reason": "大纲质量差"}, headers=once).status_code == 200  # replay
    assert credits(api, user)["balance"] == before["balance"] + settled
    assert api.post(url, json={"credits": 1, "reason": "再退"}, headers=admin).status_code == 409
    audit = api.get("/api/admin/audit", headers=admin).json()
    assert [a["data"]["credits"] for a in audit if a["action"] == "job.refund"] == [settled]
    with api.app.state.session_factory() as session:
        from autoppt_api.models import Notification

        assert session.query(Notification).filter(Notification.subject == "积分已退回").count() == 1
    statement = credits(api, user)["statement"]
    assert statement[0]["kind"] == "refund"
    ledger_ok(api)


def test_orders_need_the_current_terms_when_a_version_is_configured(tmp_path):
    settings = Settings(database_url=f"sqlite:///{(tmp_path / 'api.db').as_posix()}", storage_dir=tmp_path / "storage",
                        dev_echo_codes=True, admin_emails=("admin@example.com",), terms_version="2026-10-19")
    client = TestClient(create_app(settings))
    client.tmp = tmp_path
    admin, user = login(client, "admin@example.com"), login(client, "student@example.com")
    client.post("/api/admin/credits/adjust", json={"email": "student@example.com", "credits": 500, "memo": "t"},
                headers={**admin, "Idempotency-Key": "t1"})
    sha = upload(client, user, "论文.docx", docx_bytes())
    job = client.post("/api/jobs", json={"request": request_body(sha)}, headers=user).json()
    total = client.post(f"/api/jobs/{job['id']}/estimate", headers=user).json()["total"]
    url = f"/api/jobs/{job['id']}/submit"
    assert client.get("/api/options").json()["terms_version"] == "2026-10-19"
    assert client.post(url, json={"confirm_total": total}, headers=user).status_code == 422
    assert client.post(url, json={"confirm_total": total, "accepted_terms": "2025-01-01"}, headers=user).status_code == 422
    assert client.post(url, json={"confirm_total": total, "accepted_terms": "2026-10-19"}, headers=user).status_code == 200
    audit = client.get("/api/admin/audit", headers=admin).json()
    assert [a["data"]["version"] for a in audit if a["action"] == "job.terms_accepted"] == ["2026-10-19"]
    for name in ("terms", "privacy", "settlement"):
        page = client.get(f"/legal/{name}.md")
        assert page.status_code == 200 and "草案" in page.text
