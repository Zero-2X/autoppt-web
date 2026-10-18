"""P5 vectorization: user-supplied masters, OCR, CP-V text review, Office lock, full mock flow."""

from __future__ import annotations

import json
import threading
import time
from pathlib import Path

import pytest
from PIL import Image

from autoppt_service import vectorize
from autoppt_service.checkpoint import pending_checkpoint, submit_response
from autoppt_service.jobfiles import read_json, service_path, write_json_atomic
from autoppt_service.run_job import run_all, run_phase
from autoppt_service.status import load_job_state


def make_pptx(path: Path, slides: int = 3) -> Path:
    from pptx import Presentation

    presentation = Presentation()
    for index in range(slides):
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = f"第 {index + 1} 页"
    presentation.save(path)
    return path


def fake_renderer(pptx: Path, out_dir: Path) -> list[Path]:
    from pptx import Presentation

    out_dir.mkdir(parents=True, exist_ok=True)
    paths = []
    for index in range(1, len(Presentation(str(pptx)).slides) + 1):
        target = out_dir / f"slide-{index}.png"
        Image.new("RGB", (3200, 1800), (index * 40, 90, 120)).save(target)  # wider than a master: gets resized
        paths.append(target)
    return paths


def fake_ocr(image: Path, out_json: Path) -> dict:
    with Image.open(image) as picture:
        width, height = picture.size
    payload = {"width": width, "height": height, "lines": [
        {"id": "line-000", "text": "城 市 内 涝 预 警", "bbox": {"x": 96, "y": 54, "w": 960, "h": 108}},
        {"id": "line-001", "text": "平 均 提 前 47 分 钟", "bbox": {"x": 96, "y": 300, "w": 700, "h": 60}},
        {"id": "line-002", "text": "  ", "bbox": {"x": 0, "y": 0, "w": 1, "h": 1}},
    ]}
    write_json_atomic(out_json, payload)
    return payload


def job_with_sources(tmp_path: Path, records: list[dict]) -> Path:
    job = tmp_path / "job"
    (job / "workspace/evidence_workspace").mkdir(parents=True)
    write_json_atomic(job / "workspace/evidence_workspace/source-manifest.json", {"files": records})
    return job


def test_parse_pages_and_join_cjk():
    assert vectorize.parse_pages("", 4) == [1, 2, 3, 4]
    assert vectorize.parse_pages("1-2, 4，9", 5) == [1, 2, 4]
    with pytest.raises(vectorize.VectorizeError):
        vectorize.parse_pages("7-9", 5)
    assert vectorize.join_cjk("街 区 级 风 险 等 级 ， 平 均 提 前 47 分 钟") == "街区级风险等级，平均提前 47 分钟"
    assert vectorize.join_cjk("GNN 融 合 模 型 v2") == "GNN 融合模型 v2"


def test_prepare_masters_from_pptx_pdf_and_images(tmp_path):
    job = tmp_path / "job"
    (job / "source-materials").mkdir(parents=True)
    make_pptx(job / "source-materials/f_01__deck.pptx", 3)
    Image.new("RGB", (800, 600), "white").save(job / "source-materials/f_02__shot.png")
    Image.new("RGB", (400, 300), "white").save(job / "source-materials/f_03__scan.pdf", "PDF")
    records = [
        {"file_id": "f_01", "role": "vectorize_source", "status": "parsed", "kind": "pptx", "name": "deck.pptx",
         "path": "source-materials/f_01__deck.pptx", "sha256": "a" * 64},
        {"file_id": "f_02", "role": "vectorize_source", "status": "parsed", "kind": "image", "name": "shot.png",
         "path": "source-materials/f_02__shot.png", "sha256": "b" * 64},
        {"file_id": "f_03", "role": "vectorize_source", "status": "parsed", "kind": "pdf", "name": "scan.pdf",
         "path": "source-materials/f_03__scan.pdf", "sha256": "c" * 64},
        {"file_id": "f_04", "role": "primary_source", "status": "parsed", "kind": "text", "name": "x.md", "path": "x"},
    ]
    write_json_atomic(job / "workspace/evidence_workspace/source-manifest.json", {"files": records})

    def fake_pdf(pdf: Path, out_dir: Path):
        out_dir.mkdir(parents=True, exist_ok=True)
        target = out_dir / "page-001.png"
        Image.new("RGB", (1920, 1440), "grey").save(target)
        return [target], "fake"

    masters = vectorize.prepare_masters(job, renderer=fake_renderer, pdf_rasterizer=fake_pdf)
    slides = masters["slides"]
    assert [s["slide_id"] for s in slides] == ["S01", "S02", "S03", "S04", "S05"]
    assert all(s["provenance"] == "user_supplied_master" for s in slides)
    assert slides[0]["width"] == 1920 and slides[0]["height"] == 1080  # resized from 3200 wide
    assert slides[3]["source"] == {"file_id": "f_02", "source_file": "shot.png", "source_sha256": "b" * 64, "page": 1}
    assert slides[4]["source"]["file_id"] == "f_03"
    assert (job / "assets/slides/user-supplied-masters.json").exists()
    assert set(masters["renderers"]) == {"powerpoint", "pillow", "fake"}

    subset = vectorize.prepare_masters(job, renderer=fake_renderer, pdf_rasterizer=fake_pdf, page_spec="2,4")
    assert [s["source"]["page"] for s in subset["slides"]] == [2, 1]
    with pytest.raises(vectorize.VectorizeError, match="covers 2"):
        vectorize.prepare_masters(job, renderer=fake_renderer, pdf_rasterizer=fake_pdf, max_slides=2)


def test_real_pdf_rasterizer_if_available(tmp_path):
    pytest.importorskip("fitz")
    pdf = tmp_path / "a.pdf"
    Image.new("RGB", (400, 300), "white").save(pdf, "PDF")
    pages, name = vectorize.rasterize_pdf(pdf, tmp_path / "out", width=800)
    assert len(pages) == 1 and name
    with Image.open(pages[0]) as image:
        assert image.width == 800


def test_ocr_review_and_corrections(tmp_path):
    job = tmp_path / "job"
    (job / "assets/slides").mkdir(parents=True)
    Image.new("RGB", (1920, 1080)).save(job / "assets/slides/S01.png")
    masters = {"slides": [{"slide_id": "S01", "path": "assets/slides/S01.png", "width": 1920, "height": 1080, "sha256": "x"}]}
    review = vectorize.run_ocr(job, masters, ocr=fake_ocr)
    lines = review["slides"][0]["lines"]
    assert [l["text"] for l in lines] == ["城市内涝预警", "平均提前 47 分钟"]  # blank line dropped, CJK joined
    assert lines[0]["bbox"] == [0.05, 0.05, 0.5, 0.1]
    payload = vectorize.checkpoint_payload(job, masters, round_no=1)
    assert payload["type"] == "vectorize_text" and payload["master_sha256"] == {"S01": "x"}
    reviewed = vectorize.apply_text_review(job, {"user_reply": "已核对", "answers": {
        "text_corrections": {"S01": {"line-000": "城市内涝短临预警", "line-001": ""}},
        "added_lines": {"S01": ["本人贡献"]},
    }})
    texts = [(l["text"], l["edited"]) for l in reviewed["slides"][0]["lines"]]
    assert texts == [("城市内涝短临预警", True), ("本人贡献", True)]
    assert reviewed["exact_text"] == ["城市内涝短临预警", "本人贡献"]
    with pytest.raises(vectorize.VectorizeError):
        vectorize.apply_text_review(job, {"answers": {"text_corrections": {"S09": {"line-000": "x"}}}})


def test_office_lock_is_exclusive_and_recovers_from_a_dead_owner(tmp_path):
    lock = tmp_path / "office.lock"
    order: list[str] = []

    def worker(name: str) -> None:
        with vectorize.office_lock(lock, poll_s=0.05):
            order.append(f"{name}-in")
            time.sleep(0.2)
            order.append(f"{name}-out")

    threads = [threading.Thread(target=worker, args=(n,)) for n in ("a", "b")]
    for t in threads:
        t.start()
    for t in threads:
        t.join()
    assert order[0].endswith("-in") and order[1].endswith("-out") and order[1][0] == order[0][0]  # never interleaved
    lock.write_text("999999 2026-01-01T00:00:00Z")  # a crashed process left its lock
    with vectorize.office_lock(lock, timeout_s=2, poll_s=0.05):
        assert lock.exists()
    assert not lock.exists()


def vector_request(name: str = "旧稿.pptx", count: int = 3) -> dict:
    return {
        "schema_version": "autoppt-job-request-v1",
        "service": "vectorize",
        "inputs": {"files": [{"file_id": "f_01", "role": "vectorize_source", "name": name}]},
        "slides": {"mode": "fixed", "count": count},
        "delivery": {"tier": "native_editable", "font_policy": "office_native_yahei"},
        "vectorize": {"pages": "", "depth": "text_shapes", "keep_photos_as_images": True},
    }


def test_mock_vectorize_job_end_to_end(tmp_path):
    deck = make_pptx(tmp_path / "old.pptx", 3)
    job = tmp_path / "job"
    results = run_all(job, job_id="V1", request=vector_request(), uploads={"f_01": str(deck)}, auto_respond=True)
    assert [r["phase"] for r in results] == ["intake", "reconstruct", "package"]
    history = [h["to"] for h in load_job_state(job)["history"]]
    assert history[history.index("intake"):] == ["intake", "cp_v_text", "reconstructing", "office_qa",
                                                  "cp3_native_preview", "packaging", "delivered"]
    reviewed = read_json(job / "workspace/vectorize/reviewed-text.json")
    assert reviewed["exact_text"] and reviewed["slides"][0]["master"] == "assets/slides/S01.png"
    masters = read_json(job / "assets/slides/user-supplied-masters.json")
    assert masters["provenance"] == "user_supplied_master" and len(masters["slides"]) == 3


def test_vectorize_blocks_when_pages_exceed_the_order(tmp_path):
    deck = make_pptx(tmp_path / "old.pptx", 5)
    job = tmp_path / "job"
    result = run_phase(job, "intake", job_id="V2", request=vector_request(count=3), uploads={"f_01": str(deck)}, mock=True)
    assert result["job_state"] == "blocked"
    assert "covers 3" in load_job_state(job)["blocked_reason"]


def test_cpv_corrections_flow_through_run_job(tmp_path):
    deck = make_pptx(tmp_path / "old.pptx", 2)
    job = tmp_path / "job"
    run_phase(job, "intake", job_id="V3", request=vector_request(count=2), uploads={"f_01": str(deck)}, mock=True)
    checkpoint = pending_checkpoint(job)
    assert checkpoint["type"] == "vectorize_text"
    submit_response(job, {"checkpoint_id": checkpoint["checkpoint_id"], "decision": "approve", "responder": "user",
                          "user_reply": "改了标题", "answers": {"text_corrections": {"S01": {"line-000": "正确的标题"}}}})
    run_phase(job, "reconstruct", job_id="V3", mock=True)
    reviewed = read_json(job / "workspace/vectorize/reviewed-text.json")
    assert reviewed["slides"][0]["lines"][0]["text"] == "正确的标题"
    assert load_job_state(job)["state"] == "cp3_native_preview"


def test_master_number_check_flags_derived_figures(tmp_path):
    from autoppt_service.master_qa import check_master_numbers

    job = tmp_path / "job"
    write_json_atomic(job / "workspace/presentation_handoff/deck-plan.json", {"slides": [
        {"slide_id": "S09", "exact_text": ["从 18 分钟提高到 47 分钟"]}]})
    write_json_atomic(job / "workspace/evidence_workspace/evidence-ledger.json", {"evidence_items": [
        {"fact_or_data": "测试只有 16 场降雨，其中特大暴雨仅 2 场"}]})
    (job / "assets/slides").mkdir(parents=True)
    Image.new("RGB", (100, 60)).save(job / "assets/slides/S09.png")

    def ocr(image, out_json):
        return {"lines": [{"text": "从 18 分 钟 提 高 到 47 分 钟"}, {"text": "提 升 2.6 倍"}, {"text": "16 场"}]}

    findings = check_master_numbers(job, {"S09": job / "assets/slides/S09.png"}, ocr=ocr)
    assert findings == ["S09: master shows numbers not in the plan or evidence: 2.6"]
    assert read_json(job / "workspace/master-qa/numbers.json")["slides"][0]["unexpected"] == ["2.6"]


def test_ocr_component_splits_are_suggested_not_applied(tmp_path):
    job = tmp_path / "job"
    (job / "assets/slides").mkdir(parents=True)
    Image.new("RGB", (1920, 1080)).save(job / "assets/slides/S01.png")
    masters = {"slides": [{"slide_id": "S01", "path": "assets/slides/S01.png", "width": 1920, "height": 1080, "sha256": "x"}]}

    def ocr(image, out_json):  # what Windows OCR returned for 城市内涝频发 / 雷达外推 on the real sample
        return {"width": 1920, "height": 1080, "lines": [
            {"id": "l0", "text": "城 市 内 氵 劳 频 发", "bbox": {"x": 0, "y": 0, "w": 10, "h": 10}},
            {"id": "l1", "text": "雷 达 夕 卜 推", "bbox": {"x": 0, "y": 0, "w": 10, "h": 10}},
            {"id": "l2", "text": "研 究 背 景", "bbox": {"x": 0, "y": 0, "w": 10, "h": 10}}]}

    lines = vectorize.run_ocr(job, masters, ocr=ocr)["slides"][0]["lines"]
    assert [(l["text"], l.get("suggestion"), l.get("suspect", False)) for l in lines] == [
        ("城市内氵劳频发", "城市内涝频发", True), ("雷达夕卜推", "雷达外推", True), ("研究背景", None, False)]


def test_changed_user_supplied_master_blocks_the_rebuild(tmp_path):
    from autoppt_service.vectorize import verify_user_supplied_masters

    deck = make_pptx(tmp_path / "old.pptx", 2)
    job = tmp_path / "job"
    run_phase(job, "intake", job_id="V9", request=vector_request(count=2), uploads={"f_01": str(deck)}, mock=True)
    checkpoint = pending_checkpoint(job)
    submit_response(job, {"checkpoint_id": checkpoint["checkpoint_id"], "decision": "approve", "responder": "user",
                          "user_reply": "文字没问题"})
    assert verify_user_supplied_masters(job) == []
    Image.new("RGB", (1920, 1080), "red").save(job / "assets/slides/S02.png")  # swapped after freezing
    Image.new("RGB", (1920, 1080)).save(job / "assets/slides/S03.png")         # a page nobody uploaded
    issues = verify_user_supplied_masters(job)
    assert any("S02: master changed after it was frozen" in i for i in issues)
    assert any("S03.png" in i for i in issues)
    result = run_phase(job, "reconstruct", job_id="V9", mock=True)
    assert result["job_state"] == "blocked" and "user-supplied masters failed their gate" in load_job_state(job)["blocked_reason"]
