"""Implicit AIGC labels in delivered decks (生成合成内容标识办法, GB 45438-2025)."""

from __future__ import annotations

import zipfile

import pytest
from pptx import Presentation

from autoppt_service.aigc_label import CHANGED_PARTS, changed_parts, label_pptx, read_pptx_label


def deck(path, *, with_custom=False):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "城市内涝预警"
    presentation.save(path)
    if with_custom:  # a deck that already has custom properties keeps them
        import shutil

        shutil.copy(path, path.with_suffix(".orig"))
        from autoppt_service.aigc_label import CUSTOM_PART, CUSTOM_REL, CUSTOM_TYPE

        with zipfile.ZipFile(path.with_suffix(".orig")) as src, zipfile.ZipFile(path, "w") as dst:
            for info in src.infolist():
                data = src.read(info.filename)
                if info.filename == "[Content_Types].xml":
                    data = data.replace(b"</Types>", f'<Override PartName="/{CUSTOM_PART}" ContentType="{CUSTOM_TYPE}"/></Types>'.encode())
                if info.filename == "_rels/.rels":
                    data = data.replace(b"</Relationships>", f'<Relationship Id="rIdX" Type="{CUSTOM_REL}" Target="{CUSTOM_PART}"/></Relationships>'.encode())
                dst.writestr(info, data)
            dst.writestr(CUSTOM_PART, '<?xml version="1.0" encoding="UTF-8"?><Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties" xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes"><property fmtid="{D5CDD505-2E9C-101B-9397-08002B2CF9AE}" pid="2" name="Course"><vt:lpwstr>GIS</vt:lpwstr></property></Properties>')
    return path


@pytest.mark.parametrize("with_custom", [False, True])
def test_label_touches_only_metadata_and_survives_reopening(tmp_path, with_custom):
    path = deck(tmp_path / "deck.pptx", with_custom=with_custom)
    before = tmp_path / "before.pptx"
    before.write_bytes(path.read_bytes())
    written = label_pptx(path, "J1:deck.pptx", producer="AutoPPT")
    assert written["Label"] == "1" and written["ProduceID"] == "J1:deck.pptx"
    assert read_pptx_label(path) == written
    assert set(changed_parts(before, path)) <= set(CHANGED_PARTS)
    assert Presentation(str(path)).slides[0].shapes.title.text == "城市内涝预警"
    custom = zipfile.ZipFile(path).read("docProps/custom.xml").decode()
    assert ('name="Course"' in custom) == with_custom  # existing properties are kept
    with pytest.raises(ValueError):
        label_pptx(path, "again")  # never two labels
