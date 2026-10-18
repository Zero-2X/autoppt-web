"""Soak test: a live API, several mock workers and many users ordering, reviewing and cancelling at once.

Runs for ``--minutes`` (the launch plan asks for 48 hours before going live) and checks, every
``--check-every`` seconds, that the credit ledger balances and no phase is stuck.  At the end it waits
for open jobs to settle and writes a JSON report.  Mock workers do no Codex or Office work, so this
measures the API, the queue, leases, checkpoints and the ledger, not generation quality or cost.

    python scripts/stress_test.py --minutes 30 --users 6 --workers 3 --out stress-report.json

Use a PostgreSQL ``--database-url`` to test the production database; the default SQLite file is what
development uses and serialises writes.
"""

from __future__ import annotations

import argparse
import io
import json
import os
import random
import socket
import sys
import tempfile
import threading
import time
from pathlib import Path
from typing import Any

import httpx
import psutil

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from autoppt_service.repo_paths import ensure_autoppt_on_path  # noqa: E402

ensure_autoppt_on_path()

DESCRIPTION = "本研究面向城市内涝预警，目标是在暴雨前两小时给出街区级风险等级，并说明本人负责的模型部分。" * 6
TIERS = ("planning_package", "image_deck", "native_editable")
CHECKPOINT_STATES = {"cp1_outline", "cp2_masters", "cp3_native_preview", "cp_v_text"}
FINAL_STATES = {"delivered", "cancelled", "failed", "expired", "deleted"}


def free_port() -> int:
    with socket.socket() as sock:
        sock.bind(("127.0.0.1", 0))
        return sock.getsockname()[1]


def docx_bytes(index: int) -> bytes:
    from docx import Document

    document = Document()
    document.add_heading("研究背景", level=1)
    document.add_paragraph(f"城市内涝频发，现有预警只能给出区县级结果，平均提前 {18 + index % 20} 分钟。")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def request_body(sha: str, tier: str) -> dict[str, Any]:
    delivery = {"tier": tier, "font_policy": "office_native_yahei"} if tier == "native_editable" else {"tier": tier}
    return {"schema_version": "autoppt-job-request-v1", "service": "new_deck",
            "inputs": {"files": [{"file_id": "f_01", "role": "primary_source", "name": "论文.docx", "sha256": sha}],
                       "description": DESCRIPTION},
            "scenario": "thesis_defense", "slides": {"mode": "fixed", "count": 3},
            "cover": {"title": "压力测试"}, "delivery": delivery}


class Stats:
    def __init__(self) -> None:
        self.lock = threading.Lock()
        self.counts: dict[str, int] = {}
        self.durations: list[float] = []
        self.errors: list[str] = []

    def add(self, key: str, n: int = 1) -> None:
        with self.lock:
            self.counts[key] = self.counts.get(key, 0) + n

    def error(self, message: str) -> None:
        with self.lock:
            self.counts["errors"] = self.counts.get("errors", 0) + 1
            if len(self.errors) < 50:
                self.errors.append(message[:300])


def login(client: httpx.Client, email: str) -> dict[str, str]:
    code = client.post("/api/auth/code", json={"email": email}).json()["dev_code"]
    tokens = client.post("/api/auth/login", json={"email": email, "code": code}).json()
    return {"Authorization": f"Bearer {tokens['access_token']}"}


def upload(client: httpx.Client, headers: dict[str, str], data: bytes) -> str:
    init = client.post("/api/uploads", json={"filename": "论文.docx", "size": len(data)}, headers=headers).json()
    client.put(init["upload_url"], params={"offset": 0}, content=data).raise_for_status()
    done = client.post(f"/api/uploads/{init['upload_id']}/complete", headers=headers)
    done.raise_for_status()
    return done.json()["sha256"]


def user_loop(base: str, email: str, stop: threading.Event, stats: Stats, rng: random.Random, max_open: int) -> None:
    try:
        _user_loop(base, email, stop, stats, rng, max_open)
    except Exception as exc:  # noqa: BLE001 - a dead simulated user is a failure, not a silent gap
        stats.error(f"{email} stopped: {exc!r}")


def _user_loop(base: str, email: str, stop: threading.Event, stats: Stats, rng: random.Random, max_open: int) -> None:
    with httpx.Client(base_url=base, timeout=60) as client:
        headers = login(client, email)
        sha = upload(client, headers, docx_bytes(rng.randrange(1000)))
        open_jobs: dict[str, float] = {}
        while not stop.is_set() or open_jobs:
            try:
                if not stop.is_set() and len(open_jobs) < max_open:
                    tier = rng.choice(TIERS)
                    job = client.post("/api/jobs", json={"request": request_body(sha, tier)}, headers=headers).json()
                    total = client.post(f"/api/jobs/{job['id']}/estimate", headers=headers).json()["total"]
                    reply = client.post(f"/api/jobs/{job['id']}/submit", json={"confirm_total": total},
                                        headers={**headers, "Idempotency-Key": f"submit-{job['id']}"})
                    if reply.status_code == 200:
                        open_jobs[job["id"]] = time.monotonic()
                        stats.add(f"submitted_{tier}")
                    elif reply.status_code == 402:
                        stats.add("insufficient_credits")
                    else:
                        stats.error(f"submit {reply.status_code}: {reply.text}")
                for job_id in list(open_jobs):
                    detail = client.get(f"/api/jobs/{job_id}", headers=headers).json()
                    state = detail["state"]
                    if state in FINAL_STATES:
                        stats.add(state)
                        if state == "delivered":
                            with stats.lock:
                                stats.durations.append(time.monotonic() - open_jobs[job_id])
                        del open_jobs[job_id]
                    elif state in CHECKPOINT_STATES:
                        answer(client, headers, job_id, stats, rng)
                    elif state == "blocked":
                        stats.error(f"{job_id[:8]} blocked: {detail.get('blocked_reason')}")
                        client.post(f"/api/jobs/{job_id}/cancel", headers=headers)
            except httpx.HTTPError as exc:
                stats.error(f"{email}: {exc}")
            time.sleep(rng.uniform(0.2, 1.0))


def answer(client: httpx.Client, headers: dict[str, str], job_id: str, stats: Stats, rng: random.Random) -> None:
    checkpoint = client.get(f"/api/jobs/{job_id}/checkpoint", headers=headers)
    if checkpoint.status_code != 200:
        return
    cp = checkpoint.json()
    url = f"/api/jobs/{job_id}/checkpoint/{cp['checkpoint_id']}/response"
    roll = rng.random()
    if roll < 0.04:
        client.post(f"/api/jobs/{job_id}/cancel", headers=headers)
        stats.add("user_cancelled")
        return
    if cp["type"] == "masters" and roll < 0.15 and cp.get("round", 1) == 1:
        body = {"decision": "redo_slides", "user_reply": "第一页重做", "redo_slide_ids": ["S01"], "charge_confirmed": True}
        stats.add("redo")
    elif cp["type"] == "native_preview":
        sides = [p["path"] for p in cp["previews"] if p["kind"] == "side_by_side"]
        body = {"decision": "approve", "user_reply": "接受预览中的非文字差异", "reviewed_previews": sides}
    else:
        body = {"decision": "approve", "user_reply": "可以"}
    reply = client.post(url, json=body, headers=headers)
    if reply.status_code == 409:  # someone (a second poll) answered first
        return
    if reply.status_code != 200:
        stats.error(f"answer {cp['type']} {reply.status_code}: {reply.text}")


def worker_loop(base: str, token: str, root: Path, stop: threading.Event, stats: Stats) -> None:
    from autoppt_service.worker_client import ApiWorker

    with httpx.Client(base_url=base, timeout=120) as client:
        worker = ApiWorker(client, token, root, mock=True, heartbeat_s=30, min_free_gb=0)
        while not stop.is_set():
            try:
                if worker.run_once() is None:
                    time.sleep(0.3)
                else:
                    stats.add("phases_run")
            except Exception as exc:  # noqa: BLE001 - a soak test records and keeps going
                stats.error(f"worker {root.name}: {exc!r}")
                time.sleep(1)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("--minutes", type=float, default=10)
    parser.add_argument("--users", type=int, default=5)
    parser.add_argument("--workers", type=int, default=3)
    parser.add_argument("--max-open", type=int, default=2, help="open jobs per user")
    parser.add_argument("--check-every", type=float, default=30)
    parser.add_argument("--drain-minutes", type=float, default=10)
    parser.add_argument("--database-url")
    parser.add_argument("--work-dir", type=Path)
    parser.add_argument("--seed", type=int, default=7)
    parser.add_argument("--out", type=Path)
    args = parser.parse_args(argv)

    import uvicorn

    from autoppt_api.app import create_app
    from autoppt_api.config import Settings

    work = args.work_dir or Path(tempfile.mkdtemp(prefix="autoppt-stress-"))
    work.mkdir(parents=True, exist_ok=True)
    settings = Settings(database_url=args.database_url or f"sqlite:///{(work / 'api.db').as_posix()}",
                        storage_dir=work / "storage", dev_echo_codes=True, admin_emails=("admin@example.com",),
                        # A soak test is not about login limits: one login per simulated user for the whole run.
                        access_token_ttl_s=7 * 24 * 3600, login_codes_per_hour=10_000)
    port = free_port()
    server = uvicorn.Server(uvicorn.Config(create_app(settings), host="127.0.0.1", port=port, log_level="warning"))
    threading.Thread(target=server.run, daemon=True).start()
    base = f"http://127.0.0.1:{port}"
    for _ in range(200):
        try:
            if httpx.get(f"{base}/api/health", timeout=1).status_code == 200:
                break
        except httpx.HTTPError:
            time.sleep(0.1)

    stats = Stats()
    rng = random.Random(args.seed)
    admin_client = httpx.Client(base_url=base, timeout=60)
    admin = login(admin_client, "admin@example.com")
    stop_users, stop_workers = threading.Event(), threading.Event()
    threads = []
    for index in range(args.workers):
        token = admin_client.post("/api/admin/workers", json={"name": f"stress-{index}", "capabilities": ["all"]},
                                  headers=admin).json()["token"]
        thread = threading.Thread(target=worker_loop, args=(base, token, work / f"worker-{index}", stop_workers, stats),
                                  daemon=True)
        thread.start()
    users = []
    for index in range(args.users):
        email = f"stress{index}@example.com"
        login(admin_client, email)  # the first login creates the account
        admin_client.post("/api/admin/credits/adjust", json={"email": email, "credits": 1_000_000, "memo": "stress"},
                          headers={**admin, "Idempotency-Key": f"stress-{index}"}).raise_for_status()
        thread = threading.Thread(target=user_loop, args=(base, email, stop_users, stats, random.Random(rng.random()),
                                                          args.max_open), daemon=True)
        thread.start()
        users.append(thread)

    started = time.monotonic()
    checks = []
    deadline = started + args.minutes * 60
    while time.monotonic() < deadline:
        time.sleep(min(args.check_every, max(deadline - time.monotonic(), 0.1)))
        ledger = admin_client.get("/api/admin/ledger/check", headers=admin).json()
        health = admin_client.get("/api/admin/health", headers=admin).json()
        process = psutil.Process()
        checks.append({"t_min": round((time.monotonic() - started) / 60, 2), "ledger_ok": ledger["ok"],
                       # API, workers and users share this process: a steady rise means a leaked file or socket.
                       "open_handles": process.num_handles() if hasattr(process, "num_handles") else process.num_fds(),
                       "open_files": len(process.open_files()),
                       "ledger_issues": ledger["issues"][:5], "alerts": [a["kind"] for a in health["alerts"]],
                       "waiting": health["queue"]["waiting"], "running": health["running_phases"]})
        print(json.dumps(checks[-1] | {"counts": stats.counts}, ensure_ascii=False), flush=True)
    stop_users.set()
    for thread in users:
        thread.join(timeout=args.drain_minutes * 60)
    undrained = sum(thread.is_alive() for thread in users)
    stop_workers.set()
    try:
        final_ledger = admin_client.get("/api/admin/ledger/check", headers=admin).json()
        jobs = admin_client.get("/api/admin/jobs", headers=admin).json()
    except (httpx.HTTPError, ValueError) as exc:  # record a failing final read instead of losing the report
        stats.error(f"final checks failed: {exc!r}")
        final_ledger, jobs = {"ok": False, "issues": [repr(exc)]}, []
    durations = sorted(stats.durations)

    def pct(p: float) -> float | None:
        return round(durations[min(len(durations) - 1, int(p * len(durations)))], 1) if durations else None

    stuck_alerts = sorted({kind for check in checks for kind in check["alerts"] if kind in ("stuck_phase", "lease_expired")})
    report = {
        "minutes": args.minutes, "users": args.users, "workers": args.workers,
        "database": "sqlite" if settings.database_url.startswith("sqlite") else "other",
        "counts": stats.counts, "errors": stats.errors,
        "delivered_seconds": {"p50": pct(0.5), "p95": pct(0.95), "max": durations[-1] if durations else None},
        "final_ledger_ok": final_ledger["ok"], "final_ledger_issues": final_ledger["issues"],
        "ledger_ok_every_check": all(check["ledger_ok"] for check in checks),
        "stuck_alerts_seen": stuck_alerts,
        "users_not_drained": undrained,
        "jobs_not_final": sorted({job["state"] for job in jobs if job["state"] not in FINAL_STATES}),
        # Overlapping worker phases once left TEMP on a deleted job folder (SQLite then failed to make temp files).
        "temp_dir_exists": Path(os.environ.get("TEMP") or tempfile.gettempdir()).is_dir(),
        "checks": checks,
    }
    ok = (report["final_ledger_ok"] and report["ledger_ok_every_check"] and not stuck_alerts and not undrained
          and not report["jobs_not_final"] and not stats.counts.get("errors") and report["temp_dir_exists"])
    report["ok"] = ok
    text = json.dumps(report, ensure_ascii=False, indent=2)
    if args.out:
        args.out.write_text(text, encoding="utf-8")
    print(text if not args.out else json.dumps({k: report[k] for k in ("ok", "counts", "delivered_seconds")}, ensure_ascii=False))
    server.should_exit = True
    return 0 if ok else 1


if __name__ == "__main__":
    raise SystemExit(main())
